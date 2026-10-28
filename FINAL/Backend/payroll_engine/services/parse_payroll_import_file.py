"""Parse payroll import uploads (CSV / Excel) and map rows.

Port of ``FINAL/legacy_payroll/utils/payroll/parsePayrollImportFile.js``.

Exported (legacy name -> port):
    parsePayrollImportFile         -> parse_payroll_import_file
    mapImportRow                   -> map_import_row
    mapSalaryRevisionRow           -> map_salary_revision_row
    extractComponentAmountsFromRow -> extract_component_amounts_from_row
    extractComponentInputsFromRow  -> extract_component_inputs_from_row
    KNOWN_SALARY_REVISION_COLUMNS

Parsing reproduces the legacy libraries:

* ``text/csv`` -> a byte-level port of ``csv-parser@3.2.0`` (headers from line 1, values
  always strings, BOM kept in the first header, extra cells as ``_<index>``, an empty line
  yields ``{}``).
* Excel -> ``xlsx@0.18.5`` ``sheet_to_json(sheet, { defval: '' })`` on the first sheet via
  openpyxl: header row = first row of the sheet's dimension, empty header cells become
  ``__EMPTY``, ``__EMPTY_1``..., duplicate headers get ``_1``, ``_2``...; values keep their
  raw type (numbers, booleans, strings; dates as Excel serial numbers like SheetJS without
  ``cellDates``); missing cells -> ``''``; rows with no value are skipped.
* JS object key order is kept (integer-like keys first, like a JS object).

``file`` is a Django ``UploadedFile`` or a dict ``{"buffer", "mimetype", "originalname"}``
(the multer shape). Sync (legacy was async). Errors raise ``ValueError`` with the legacy
messages.
"""

import io
import math
import re
import zipfile
from datetime import date, datetime, time, timedelta

from .jsmath import JS_WS_CLASS, UNDEFINED, is_nullish, js_str, js_trim, to_number

# ------------------------------------------------------------------ JS object emulation

_INDEX_KEY = re.compile(r"^(?:0|[1-9][0-9]*)$")


def js_object(pairs):
    """Build a dict with JS property-assignment semantics: a repeated key keeps its first
    position (value overwritten) and array-index-like keys come first, ascending."""
    out = {}
    for k, v in pairs:
        out[k] = v
    idx = [k for k in out if _INDEX_KEY.match(k) and int(k) < 4294967295]
    if not idx:
        return out
    ordered = {k: out[k] for k in sorted(idx, key=int)}
    ordered.update((k, v) for k, v in out.items() if k not in ordered)
    return ordered


def _is_finite_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


# ------------------------------------------------------------------ csv-parser port

_CR, _NL, _QUOTE, _COMMA = 0x0D, 0x0A, 0x22, 0x2C


def _parse_csv(data):
    """``Readable.from(buffer).pipe(csv())`` — single chunk, default options."""
    buf = bytearray(data)
    state = {"escaped": False, "first": True, "previousEnd": 0, "quoted": False}
    opts = {"newline": _NL}
    headers = []
    rows = []

    def parse_cell(start, end):
        if start < len(buf) and buf[start] == _QUOTE and end - 1 >= 0 and buf[end - 1] == _QUOTE:
            start += 1
            end -= 1
        y = start
        i = start
        while i < end:
            if buf[i] == _QUOTE and i + 1 < end and buf[i + 1] == _QUOTE:
                i += 1
            if y != i:
                buf[y] = buf[i]
            y += 1
            i += 1
        return bytes(buf[start:y]).decode("utf-8", errors="replace")

    def parse_line(start, end):
        nonlocal headers
        end -= 1
        if len(buf) and end - 1 >= 0 and buf[end - 1] == _CR:
            end -= 1
        cells = []
        is_quoted = False
        offset = start
        i = start
        while i < end:
            c = buf[i]
            starting = (not is_quoted) and c == _QUOTE
            ending = is_quoted and c == _QUOTE and i + 1 <= end and i + 1 < len(buf) \
                and buf[i + 1] == _COMMA
            escape = is_quoted and c == _QUOTE and i + 1 < end and buf[i + 1] == _QUOTE
            if starting or ending:
                is_quoted = not is_quoted
                i += 1
                continue
            if escape:
                i += 2
                continue
            if c == _COMMA and not is_quoted:
                cells.append(parse_cell(offset, i))
                offset = i + 1
            i += 1
        if offset < end:
            cells.append(parse_cell(offset, end))
        if end - 1 >= 0 and buf[end - 1] == _COMMA:
            cells.append("")
        if state["first"]:
            state["first"] = False
            headers = list(cells)
            return
        pairs = []
        for index, cell in enumerate(cells):
            header = headers[index] if index < len(headers) else UNDEFINED
            pairs.append((header if header is not UNDEFINED else f"_{index}", cell))
        rows.append(js_object(pairs))

    length = len(buf)
    i = 0
    while i < length:
        chr_ = buf[i]
        nxt = buf[i + 1] if i + 1 < length else None
        if not state["escaped"] and chr_ == _QUOTE and nxt == _QUOTE and i != 0:
            state["escaped"] = True
            i += 1
            continue
        if chr_ == _QUOTE:
            if state["escaped"]:
                state["escaped"] = False
            else:
                state["quoted"] = not state["quoted"]
            i += 1
            continue
        if not state["quoted"]:
            if state["first"]:
                if chr_ == _NL:
                    opts["newline"] = _NL
                elif chr_ == _CR and nxt != _NL:
                    opts["newline"] = _CR
            if chr_ == opts["newline"]:
                parse_line(state["previousEnd"], i + 1)
                state["previousEnd"] = i + 1
        i += 1
    if state["previousEnd"] != length and not state["escaped"] and length > 0:
        parse_line(state["previousEnd"], length + 1)
    return rows


# ------------------------------------------------------------------ SheetJS port


def _general_format(v):
    """SheetJS General number format (approximation for header cells)."""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int) or (isinstance(v, float) and v.is_integer() and abs(v) < 1e21):
        return js_str(v)
    text = f"{v:.10g}"
    if "e" in text:
        mant, exp = text.split("e")
        exp_n = int(exp)
        return f"{mant}E{'+' if exp_n >= 0 else '-'}{abs(exp_n):02d}"
    return text


def _to_serial(value, epoch):
    from openpyxl.utils.datetime import to_excel

    if isinstance(value, (datetime, date, time, timedelta)):
        serial = to_excel(value, epoch)
        return float(serial) if not float(serial).is_integer() else int(serial)
    return value


def _cell_raw(cell, epoch):
    """(present, value) like SheetJS ``make_json_row`` for one openpyxl read-only cell."""
    value = getattr(cell, "value", None)
    dtype = getattr(cell, "data_type", None)
    if value is None:
        if dtype in ("s", "inlineStr", "str"):  # empty string cell: SheetJS keeps v = ''
            return True, ""
        return False, None
    if dtype == "e":
        return True, (None if str(value) == "#NULL!" else UNDEFINED)
    if dtype == "f":  # formula without cached value
        return False, None
    return True, _to_serial(value, epoch)


def _parse_xlsx(data):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if not wb.sheetnames:
            return []
        ws = wb[wb.sheetnames[0]]
        try:
            min_row, min_col, max_row, max_col = ws.min_row, ws.min_column, ws.max_row, ws.max_column
        except Exception:  # noqa: BLE001 - unsized sheet
            min_row = min_col = max_row = max_col = None
        if not max_row or not max_col:
            return []
        epoch = wb.epoch
        grid = list(ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col,
                                 max_col=max_col))
    finally:
        wb.close()
    if not grid:
        return []
    ncols = max_col - min_col + 1
    header_cells = list(grid[0]) + [None] * (ncols - len(grid[0]))
    hdr = []
    for cell in header_cells[:ncols]:
        present, val = _cell_raw(cell, epoch) if cell is not None else (False, None)
        if not present or val is UNDEFINED or val is None:
            v = "__EMPTY"
        elif isinstance(val, str):
            v = val
        else:
            v = _general_format(val)
        vv = v
        counter = 0
        for existing in hdr:
            if existing == vv:
                counter += 1
                vv = f"{v}_{counter}"
        hdr.append(vv)
    rows = []
    for raw_row in grid[1:]:
        cells = list(raw_row) + [None] * (ncols - len(raw_row))
        pairs = []
        isempty = True
        for c, cell in enumerate(cells[:ncols]):
            present, val = _cell_raw(cell, epoch) if cell is not None else (False, None)
            if not present:
                pairs.append((hdr[c], ""))
                continue
            if val is None:  # #NULL!
                pairs.append((hdr[c], None))
                continue
            if val is UNDEFINED:
                pairs.append((hdr[c], ""))
                continue
            pairs.append((hdr[c], val))
            isempty = False
        if not isempty:
            rows.append(js_object(pairs))
    return rows


_SHEETJS_NUM = re.compile(r"^[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?$")


def _sheetjs_csv_value(text):
    """Approximation of SheetJS CSV cell typing (numbers / TRUE / FALSE)."""
    if text == "":
        return None
    if text.startswith("="):
        return ""
    if text == "TRUE":
        return True
    if text == "FALSE":
        return False
    t = text.strip()
    pct = t.endswith("%")
    core = t[:-1].strip() if pct else t
    neg = core.startswith("(") and core.endswith(")")
    if neg:
        core = core[1:-1]
    core = core.replace(",", "").replace("$", "")
    if core and _SHEETJS_NUM.match(core):
        n = float(core)
        n = -n if neg else n
        if pct:
            n = n / 100
        if n == 0:
            n = 0
        return int(n) if float(n).is_integer() and abs(n) < 1e21 else n
    return text


def _parse_text_via_sheetjs(data):
    """SheetJS ``read`` of a CSV-like text upload (mimetype said Excel)."""
    rows = _parse_csv(data)
    out = []
    for row in rows:
        pairs = []
        isempty = True
        for k, v in row.items():
            val = _sheetjs_csv_value(v)
            if val is None:
                pairs.append((k, ""))
            else:
                pairs.append((k, val))
                isempty = False
        if not isempty:
            out.append(js_object(pairs))
    return out


def _file_parts(file):
    if isinstance(file, dict):
        return file.get("buffer"), file.get("mimetype") or "", file.get("originalname")
    if file is None:
        return None, "", None
    try:
        file.seek(0)
    except Exception:  # noqa: BLE001
        pass
    data = file.read()
    return data, getattr(file, "content_type", "") or "", getattr(file, "name", None)


def parse_payroll_import_file(file):
    data, mimetype, name = _file_parts(file)
    if data is None:
        raise ValueError("No file uploaded")
    if mimetype == "text/csv":
        return _parse_csv(data)
    if ("excel" in mimetype or "spreadsheet" in mimetype
            or (name or "").endswith(".xlsx") or (name or "").endswith(".xls")):
        if data[:4] == b"PK\x03\x04" and zipfile.is_zipfile(io.BytesIO(data)):
            return _parse_xlsx(data)
        if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ValueError("Legacy .xls workbooks are not supported. Upload CSV or Excel (.xlsx)")
        return _parse_text_via_sheetjs(data)
    raise ValueError("Invalid file type. Upload CSV or Excel (.xlsx)")


# ------------------------------------------------------------------ row mapping

_JS_WS_RUN = re.compile(JS_WS_CLASS + "+")
_MONEY_CHARS = re.compile("[₹$€," + JS_WS_CLASS[1:-1] + "]")


def normalize_import_row_keys(row=UNDEFINED):
    row = {} if row is UNDEFINED or row is None else row
    pairs = []
    for key, value in row.items():
        k = _JS_WS_RUN.sub("_", js_trim(js_str(key if key else "")).lower())
        pairs.append((k, value))
    return js_object(pairs)


def pick_row_value(row, keys=()):
    for key in keys:
        val = row.get(key, UNDEFINED)
        if not is_nullish(val) and js_trim(js_str(val)) != "":
            return js_trim(js_str(val))
    return ""


def pick_raw_row_value(row, keys=()):
    for key in keys:
        if key not in row:
            continue
        val = row[key]
        if is_nullish(val):
            continue
        if isinstance(val, str) and js_trim(val) == "":
            continue
        return val
    return UNDEFINED


def map_import_row(raw_row, row_number):
    row = normalize_import_row_keys(raw_row)
    employee_code = pick_row_value(row, ["employee_code", "employeecode", "unique_id", "uniqueid",
                                         "employee_id", "employeeid"])
    category = pick_row_value(row, ["component_category", "componentcategory", "category"]).lower()
    component_code = pick_row_value(row, ["component_code", "componentcode", "code"])
    component_name = pick_row_value(row, ["component_name", "componentname", "name"])
    amount_raw = pick_row_value(row, ["amount", "overrideamount", "value"])
    amount = to_number(amount_raw)
    remarks = pick_row_value(row, ["remarks", "remark", "notes"])
    if category == "earning":
        category = "earnings"
    if category == "deduction":
        category = "deductions"
    if category == "benefit":
        category = "benefits"
    return {
        "rowNumber": row_number,
        "employeeCode": employee_code,
        "category": category,
        "componentCode": component_code,
        "componentName": component_name,
        "amount": amount if math.isfinite(amount) else float("nan"),
        "remarks": remarks,
    }


def parse_optional_import_number(value=UNDEFINED):
    if is_nullish(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value if math.isfinite(value) else float("nan")
    raw = js_trim(js_str(value))
    if not raw:
        return None
    cleaned = _MONEY_CHARS.sub("", raw)
    if not cleaned:
        return None
    n = to_number(cleaned)
    return n if math.isfinite(n) else float("nan")


def is_basic_import_header(key=""):
    k = re.sub(r"[^a-z0-9]", "", js_trim(js_str(key if key else "")).lower())
    if not k:
        return False
    return k in ("basic", "base", "basesalary", "basicsalary", "basicpay", "newbasic",
                 "newbasicsalary", "newbasicpay")


def pick_basic_salary_from_row(row=UNDEFINED):
    row = {} if row is UNDEFINED or row is None else row
    known = parse_optional_import_number(pick_raw_row_value(row, [
        "basic_salary", "basicsalary", "basic", "base_salary", "basesalary", "base", "basic_pay",
        "basicpay", "new_basic", "newbasic"]))
    if known is not None and _is_finite_number(known):
        return known
    for key, value in row.items():
        if not is_basic_import_header(key):
            continue
        n = parse_optional_import_number(value)
        if n is not None and _is_finite_number(n):
            return n
    return None


KNOWN_SALARY_REVISION_COLUMNS = {
    "employee_code", "employeecode", "unique_id", "uniqueid", "employee_id", "employeeid",
    "revision_type", "revisiontype", "type", "gross_salary", "grosssalary", "gross", "ctc",
    "annual_ctc", "annualctc", "state", "pt_state", "ptstate", "tax_regime", "taxregime",
    "payment_frequency", "paymentfrequency", "frequency", "effective_from", "effectivefrom",
    "salary_effective_from", "salaryeffectivefrom", "arrears_applicable", "arrearsapplicable",
    "arrears_from", "arrearsfrom", "arrears_from_date", "arrearsfromdate", "auto_approve",
    "autoapprove", "auto_implement", "autoimplement", "remarks", "remark", "notes",
    "base_salary", "basesalary", "basic_salary", "basicsalary", "basic", "basic_pay", "basicpay",
    "new_basic", "newbasic",
    "epf", "epf_applicable", "epfapplicable", "pf", "pf_on_actuals", "pfonactuals",
    "contribute_pf_on_actuals", "contributepfonactuals", "epf_on_actuals", "epfonactuals",
    "pf_actuals", "pfactuals", "esi", "esic", "esi_applicable", "esicapplicable",
    "esiapplicable", "esic_applicable", "lwf", "lwf_applicable", "lwfapplicable", "pt",
    "professional_tax", "professionaltax", "pt_applicable", "ptapplicable",
}


def extract_component_inputs_from_row(normalized_row=UNDEFINED):
    normalized_row = {} if normalized_row is UNDEFINED or normalized_row is None else normalized_row
    amounts = []
    blank = []
    for key, value in normalized_row.items():
        if key in KNOWN_SALARY_REVISION_COLUMNS or is_basic_import_header(key):
            continue
        if isinstance(value, bool):
            continue
        if is_nullish(value) or js_trim(js_str(value)) == "":
            blank.append(key)
            continue
        n = to_number(value)
        if not math.isfinite(n):
            continue
        amounts.append((key, n))
    return {"componentAmounts": js_object(amounts), "blankComponentHeaders": blank}


def extract_component_amounts_from_row(normalized_row=UNDEFINED):
    return extract_component_inputs_from_row(normalized_row)["componentAmounts"]


def parse_optional_import_boolean(value=UNDEFINED):
    if is_nullish(value):
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if value == 1:
            return True
        if value == 0:
            return False
        return None
    raw = js_trim(js_str(value)).lower()
    if not raw:
        return None
    if raw in ("true", "yes", "y", "1"):
        return True
    if raw in ("false", "no", "n", "0"):
        return False
    return None


def _to_bool_cell(raw):
    flag = parse_optional_import_boolean(raw)
    if flag is None:
        return ""
    return "true" if flag else "false"


def map_salary_revision_row(raw_row, row_number):
    row = normalize_import_row_keys(raw_row)
    employee_code = pick_row_value(row, ["employee_code", "employeecode", "unique_id", "uniqueid",
                                         "employee_id", "employeeid"])
    revision_type_raw = pick_row_value(row, ["revision_type", "revisiontype", "type"])
    gross_raw = pick_row_value(row, ["gross_salary", "grosssalary", "gross"])
    basic_raw = pick_basic_salary_from_row(row)
    ctc_raw = pick_row_value(row, ["ctc", "annual_ctc", "annualctc"])
    state = pick_row_value(row, ["state", "pt_state", "ptstate"])
    tax_regime_raw = pick_row_value(row, ["tax_regime", "taxregime"])
    freq_raw = pick_row_value(row, ["payment_frequency", "paymentfrequency", "frequency"])
    effective_from = pick_row_value(row, ["effective_from", "effectivefrom", "salary_effective_from",
                                          "salaryeffectivefrom"])
    auto_approve_raw = pick_raw_row_value(row, ["auto_approve", "autoapprove"])
    auto_implement_raw = pick_raw_row_value(row, ["auto_implement", "autoimplement"])
    remarks = pick_row_value(row, ["remarks", "remark", "notes"])
    epf_raw = pick_raw_row_value(row, ["epf", "epf_applicable", "epfapplicable", "pf"])
    pf_actuals_raw = pick_raw_row_value(row, [
        "pf_on_actuals", "pfonactuals", "contribute_pf_on_actuals", "contributepfonactuals",
        "epf_on_actuals", "epfonactuals", "pf_actuals", "pfactuals"])
    esi_raw = pick_raw_row_value(row, ["esic", "esi", "esi_applicable", "esic_applicable",
                                       "esiapplicable", "esicapplicable"])
    lwf_raw = pick_raw_row_value(row, ["lwf", "lwf_applicable", "lwfapplicable"])
    pt_raw = pick_raw_row_value(row, ["professionaltax", "professional_tax", "pt", "pt_applicable",
                                      "ptapplicable"])
    arrears_app_raw = pick_raw_row_value(row, ["arrears_applicable", "arrearsapplicable"])
    arrears_from = pick_row_value(row, ["arrears_from", "arrearsfrom", "arrears_from_date",
                                        "arrearsfromdate"])

    gross = to_number(gross_raw)
    base_salary = parse_optional_import_number(basic_raw if basic_raw is not None else None)
    ctc_input = parse_optional_import_number(ctc_raw)

    tax_regime = tax_regime_raw
    if tax_regime:
        n = tax_regime.lower()
        if n == "old":
            tax_regime = "Old"
        elif n == "new":
            tax_regime = "New"
    freq = freq_raw
    if freq:
        n = freq.lower()
        if n == "monthly":
            freq = "Monthly"
        elif n in ("bi-weekly", "biweekly"):
            freq = "Bi-Weekly"
        elif n == "weekly":
            freq = "Weekly"

    comp = extract_component_inputs_from_row(row)
    flags = {}
    for key, raw in (("epfApplicable", epf_raw), ("pfOnActuals", pf_actuals_raw),
                     ("esiApplicable", esi_raw), ("lwfApplicable", lwf_raw),
                     ("ptApplicable", pt_raw)):
        f = parse_optional_import_boolean(raw)
        if f is not None:
            flags[key] = f

    return {
        "rowNumber": row_number,
        "employeeCode": employee_code,
        "revisionType": revision_type_raw,
        "grossSalary": gross if math.isfinite(gross) else float("nan"),
        "baseSalary": base_salary,
        "ctcInput": ctc_input,
        "state": state,
        "taxRegime": tax_regime,
        "paymentFrequency": freq,
        "effectiveFrom": effective_from,
        "arrearsApplicable": _to_bool_cell(arrears_app_raw),
        "arrearsFromDate": arrears_from,
        "autoApprove": _to_bool_cell(auto_approve_raw),
        "autoImplement": _to_bool_cell(auto_implement_raw),
        "remarks": remarks,
        "componentAmounts": comp["componentAmounts"],
        "blankComponentHeaders": comp["blankComponentHeaders"],
        "statutoryFlags": flags,
    }
