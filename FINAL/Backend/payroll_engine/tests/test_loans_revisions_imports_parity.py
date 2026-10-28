"""Differential (parity) tests vs the legacy JS for the loan / VPF / arrears / salary
revision / payroll import ports.

Each case runs the same JSON input through the legacy JS function (one Node process per
group, ``legacy_js.run_js_batch``) and the Python port; the JSON text of both results must be
identical (``JSON.stringify`` semantics: dates as ISO strings, NaN -> null). Module-private
legacy helpers (and the ones needing a callback / Mongoose document) are reached through a
generated JS shim that compiles the legacy file with an extra export. Skipped when Node or
the legacy copy is unavailable.
"""

import base64
import io
import itertools
import json
import os
import tempfile
import unittest
from datetime import date, datetime

from django.test import SimpleTestCase

from payroll_engine.services import (loan_helper, loan_perquisite_helper, one_time_import_helper,
                                     parse_payroll_import_file as pf, salary_revision_import_helper
                                     as sri, salary_structure_import_helper as ssi)
from payroll_engine.services.jsmath import UNDEFINED
from payroll_engine.services.loan_helper import js_json
from payroll_engine.views import (arrears_management, loan_management, payroll_import,
                                  salary_revision, vpf_management)

from .legacy_js import LEGACY_DIR, SKIP_REASON, node_available, run_js_batch
from .test_calc_parity import (BONUS_BASE, COMPONENT_SETS, LWF_BASE, OVERRIDE_SETS, PT_BASE,
                               py_decode)

# ------------------------------------------------------------------ JS shim

SHIM = r"""
const fs = require('fs'); const path = require('path'); const Module = require('module');
const root = %(root)s;
function load(rel, names) {
  const file = path.join(root, rel);
  const src = fs.readFileSync(file, 'utf8') + '\nmodule.exports.__t = {' + names.join(',') + '};';
  const m = new Module(file, module); m.filename = file;
  m.paths = Module._nodeModulePaths(path.dirname(file));
  m._compile(src, file);
  return m.exports.__t;
}
const LM = load('controller/payroll/loanManagement.js', ['parsePerquisiteRate', 'parsePositiveAmount',
  'parseDateInput', 'formatDisplayDate', 'buildEmiStartPeriod', 'parseEmiStartPeriodInput',
  'computeRepaymentSummary', 'computeLoanClosingDate', 'enrichLoanRecord', 'buildRepaymentMessage',
  'resolveEmployeeDisplayName', 'normalizePaymentMode']);
const VPF = load('controller/payroll/vpfManagement.js', ['normalizeVpf', 'sanitizeVpfForStatus',
  'formatVpfResponse', 'mergePayrollPreferencesWithVpf', 'validateStartPayload',
  'calculateMonthlyVpfAmount', 'isVpfEffectiveForPayroll', 'toVpfAdminListItem']);
const AR = load('controller/payroll/arrearsManagement.js', ['normalizeMonthValue']);
const SR = load('controller/payroll/salaryRevision.js', ['mapApprovalForResponse']);
const PI = load('controller/payroll/payrollImport.js', ['normalizeMonthInput', 'normalizeYearInput']);
const LH = load('utils/payroll/loanHelper.js', ['getPayrollPeriodIndex', 'getDatePeriodIndex',
  'getPayrollRepaymentDate']);
const LP = load('utils/payroll/loanPerquisiteHelper.js', ['resolvePerquisiteMonthCount']);
const PF = require(path.join(root, 'utils/payroll/parsePayrollImportFile.js'));
const PFI = load('utils/payroll/parsePayrollImportFile.js', ['normalizeImportRowKeys',
  'pickRowValue', 'pickRawRowValue', 'parseOptionalImportNumber', 'isBasicImportHeader',
  'pickBasicSalaryFromRow', 'parseOptionalImportBoolean']);
const OT = require(path.join(root, 'utils/payroll/oneTimeImportHelper.js'));
const SRI = load('utils/payroll/salaryRevisionImportHelper.js', ['normalizeExcelMatchKey',
  'isBalancingAllowanceComponent', 'buildOrgComponentCatalog', 'componentDedupeKey',
  'isSameComponentOverride', 'toDefaultComponentsArray', 'parseOptionalNumber',
  'parseEffectiveDate', 'parseBoolean', 'mergeEmployeeStatutoryForCalc', 'normalizeRevisionType',
  'calculateArrears', 'resolvePlannedStatus', 'applyPfOnActualsToStatutoryInfo']);
const SRIX = require(path.join(root, 'utils/payroll/salaryRevisionImportHelper.js'));
const SSI = load('utils/payroll/salaryStructureImportHelper.js', ['buildRevisionSalaryComponents']);

// payrollRun.js getMonthDateRange (callback the run engine passes to isVpfEffectiveForPayroll)
function getMonthDateRange(year, monthName) {
  const monthIndex = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'].indexOf(monthName);
  if (monthIndex === -1) return null;
  return { startDate: new Date(year, monthIndex, 1, 0, 0, 0, 0),
           endDate: new Date(year, monthIndex + 1, 0, 23, 59, 59, 999) };
}
const withMark = (fn) => (emp, ...rest) => {
  emp.markModified = () => {}; const ret = fn(emp, ...rest); delete emp.markModified;
  return { ret, emp };
};
module.exports = {
  ...Object.fromEntries(Object.entries(LM).map(([k, v]) => ['lm_' + k, v])),
  ...Object.fromEntries(Object.entries(VPF).map(([k, v]) => ['vpf_' + k, v])),
  ...Object.fromEntries(Object.entries(LH).map(([k, v]) => ['lh_' + k, v])),
  ...Object.fromEntries(Object.entries(PFI).map(([k, v]) => ['pf_' + k, v])),
  ...Object.fromEntries(Object.entries(SRI).map(([k, v]) => ['sri_' + k, v])),
  lp_resolvePerquisiteMonthCount: LP.resolvePerquisiteMonthCount,
  ssi_buildRevisionSalaryComponents: SSI.buildRevisionSalaryComponents,
  vpf_isVpfEffective: (vpf, month, year) => VPF.isVpfEffectiveForPayroll(vpf, month, year, getMonthDateRange),
  ar_normalizeMonthValue: AR.normalizeMonthValue,
  sr_mapApprovalForResponse: SR.mapApprovalForResponse,
  pi_normalizeMonthInput: PI.normalizeMonthInput,
  pi_normalizeYearInput: PI.normalizeYearInput,
  parseFile: (f) => PF.parsePayrollImportFile({ buffer: Buffer.from(f.b64, 'base64'),
    mimetype: f.mimetype, originalname: f.originalname }),
  parseAndMapOneTime: async (f) => (await PF.parsePayrollImportFile({ buffer: Buffer.from(f.b64, 'base64'),
    mimetype: f.mimetype, originalname: f.originalname })).map((r, i) => PF.mapImportRow(r, i + 1)),
  parseAndMapRevision: async (f) => (await PF.parsePayrollImportFile({ buffer: Buffer.from(f.b64, 'base64'),
    mimetype: f.mimetype, originalname: f.originalname })).map((r, i) => PF.mapSalaryRevisionRow(r, i + 1)),
  groupOneTime: (rows) => Object.fromEntries(OT.groupOneTimeRowsByEmployee(rows)),
  applyFlags: withMark(SRIX.applyStatutoryFlagsToEmployee),
  applyState: withMark(SRIX.applyStateToEmployee),
};
"""

_SHIM_PATH = None


def shim_module():
    global _SHIM_PATH
    if _SHIM_PATH is None:
        d = tempfile.mkdtemp(prefix="pe_parity_shim_")
        path = os.path.join(d, "shim.js")
        with open(path, "w") as fh:
            fh.write(SHIM % {"root": json.dumps(LEGACY_DIR)})
        _SHIM_PATH = os.path.relpath(path, LEGACY_DIR)
    return _SHIM_PATH


# ------------------------------------------------------------------ case registry

CASES = {}
SHIM_MOD = object()


class Case:
    def __init__(self, module, fn, py, args, obj=False, mutated=None):
        self.module, self.fn, self.py, self.args = module, fn, py, args
        self.obj, self.mutated = obj, mutated

    def js_call(self):
        module = shim_module() if self.module is SHIM_MOD else self.module
        call = {"module": module, "fn": self.fn, "args": self.args}
        if self.mutated:
            call["mutated"] = self.mutated
        return call

    def label(self):
        return f"{self.fn}({json.dumps(self.args, default=str)[:300]})"


def C(group, module, fn, py, *args, mutated=None):
    CASES.setdefault(group, []).append(Case(module, fn, py, list(args), mutated=mutated))


def S(group, fn, py, *args, mutated=None):
    C(group, SHIM_MOD, fn, py, *args, mutated=mutated)


def O(group, module, fn, py, obj):
    CASES.setdefault(group, []).append(Case(module, fn, py, [obj], obj=True))


def _snake(name):
    import re

    s = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", name)
    return re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", s).lower()


def enc(v):
    if v is UNDEFINED:
        return {"$undefined": True}
    return js_json(v)


def run_py(case):
    args = py_decode(case.args)
    try:
        if case.obj:
            ret = case.py(**{_snake(k): v for k, v in args[0].items()})
        else:
            ret = case.py(*args)
        if case.mutated:
            return {"ret": enc(ret), "args": {str(i): enc(args[i]) for i in case.mutated}}
        return enc(ret)
    except Exception as exc:  # noqa: BLE001
        return {"$error": f"{type(exc).__name__}: {exc}"}


def is_error(v):
    return isinstance(v, dict) and list(v) == ["$error"]


def txt(v):
    return json.dumps(v, ensure_ascii=False, separators=(",", ":"))


LH = "utils/payroll/loanHelper"
LP = "utils/payroll/loanPerquisiteHelper"
OT = "utils/payroll/oneTimeImportHelper"
SRI = "utils/payroll/salaryRevisionImportHelper"
SSI = "utils/payroll/salaryStructureImportHelper"
PF = "utils/payroll/parsePayrollImportFile"
VPF = "controller/payroll/vpfManagement"

# ================================================================== loan helper
MONTH_INPUTS = ["Jan", "jan", "JANUARY", "september", "Sept", 3, "3", "12", "13", "0", "1.0", " 4 ",
                "", None, "Foo", 4.5]
for m in MONTH_INPUTS:
    C("loan", LH, "normalizePayrollMonth", loan_helper.normalize_payroll_month, m)
for m, y in itertools.product(["Mar", "april", "13", None, "x"], [2025, "2025", None, "abc"]):
    S("loan", "lh_getPayrollPeriodIndex", loan_helper.get_payroll_period_index, m, y)
for d in ["2025-04-01", "2025-03-31T18:30:00.000Z", "2025-04-01T00:00:00+05:30", None, "garbage",
          "04/15/2025", 0]:
    S("loan", "lh_getDatePeriodIndex", loan_helper.get_date_period_index, d)
for o, i in itertools.product([0, 100, 999.5, 1000, None, -5, "500"], [0, 100, 333, None, 1000, "100"]):
    C("loan", LH, "computeRemainingInstallments", loan_helper.compute_remaining_installments, o, i)
LOANS = [
    {"emiStartDate": "2025-04-01T00:00:00.000Z", "outstandingBalance": 50000, "status": "active",
     "installmentAmount": 5000, "loanAmount": 60000},
    {"emiStartDate": "2025-11-01T00:00:00.000Z", "outstandingBalance": 4999, "status": "active",
     "installmentAmount": 5000, "loanAmount": 60000},
    {"emiStartDate": "2025-12-01T00:00:00.000Z", "outstandingBalance": 0, "status": "active",
     "installmentAmount": 3000, "loanAmount": 9000},
    {"emiStartDate": "2025-04-01", "outstandingBalance": 12345.5, "status": "closed",
     "installmentAmount": 1000, "loanAmount": 20000},
    {"emiStartDate": "bad", "outstandingBalance": 100, "installmentAmount": 30, "loanAmount": 100},
    {"emiStartDate": None, "outstandingBalance": 100, "installmentAmount": 30, "loanAmount": 100},
    {"outstandingBalance": 7001, "installmentAmount": 2000, "loanAmount": 10000},
    {"emiStartDate": "2024-12-01T00:00:00.000Z", "outstandingBalance": 1000, "installmentAmount": 0,
     "loanAmount": 1000},
]
for loan, n in itertools.product(LOANS, [0, 1, 13, "2", None, -3]):
    C("loan", LH, "computeNextInstallmentPeriod", loan_helper.compute_next_installment_period, loan, n)
C("loan", LH, "computeNextInstallmentPeriod", loan_helper.compute_next_installment_period, None)
for loan in LOANS:
    C("loan", LH, "applyScheduleAfterRepayment", loan_helper.apply_schedule_after_repayment,
      dict(loan), mutated=[0])
    C("loan", LH, "computePayrollEmiAmount", loan_helper.compute_payroll_emi_amount, loan)
for start, m, y in itertools.product(["2025-04-01T00:00:00.000Z", "2025-03-31T18:30:00.000Z", None, "x"],
                                     ["Mar", "Apr", "May", "13", "June"], [2025, "2025", 2024, "x"]):
    C("loan", LH, "isEmiDueForPayrollMonth", loan_helper.is_emi_due_for_payroll_month, start, m, y)
for m, y in [("Apr", 2025), ("Dec", 2025), ("Jan", "2026"), ("3", 2025)]:
    S("loan", "lh_getPayrollRepaymentDate", loan_helper.get_payroll_repayment_date, m, y)

# ================================================================== loan perquisite
PQ_LOANS = [{"loanAmount": 500000, "outstandingBalance": 400000, "emiStartDate": "2025-04-01"},
            {"loanAmount": 20000, "outstandingBalance": 20000},
            {"loanAmount": 20001, "outstandingBalance": 15000.5},
            {"loanAmount": 900000, "outstandingBalance": 700000, "exemptFromPerquisite": True},
            {"loanAmount": 900000, "outstandingBalance": 0}, {}, {"loanAmount": "50000",
                                                                  "outstandingBalance": "40000"}]
PQ_TYPES = [{"loanName": "Housing", "perquisiteRate": 8.5}, {"loanName": " Medical Loan ",
                                                              "perquisiteRate": 9},
            {"loanName": "Car", "perquisiteRate": "7.2"}, {"loanName": "Car", "perquisiteRate": 0},
            {}, {"perquisiteRate": "abc"}]
for loan, lt in itertools.product(PQ_LOANS, PQ_TYPES):
    C("perq", LP, "isLoanPerquisiteExempt", loan_perquisite_helper.is_loan_perquisite_exempt, loan, lt)
    for charged in [0, 3, "2.5", 10, None]:
        C("perq", LP, "calculateMonthlyLoanPerquisiteForLoan",
          loan_perquisite_helper.calculate_monthly_loan_perquisite_for_loan, loan, lt, charged)
    for months in [12, 7, "5", None, "", -1, "x", 0]:
        C("perq", LP, "calculateAnnualLoanPerquisiteForLoan",
          loan_perquisite_helper.calculate_annual_loan_perquisite_for_loan, loan, lt, months)
for months in [12, 0, "3", None, "", -2, "abc", 4.5]:
    S("perq", "lp_resolvePerquisiteMonthCount", loan_perquisite_helper.resolve_perquisite_month_count,
      months)
for p, rec in itertools.product([0, 1000, 2833, "1500", None, -5],
                                [None, {}, {"taxableIncome": 1200000, "taxPayable": {"total": 150000}},
                                 {"taxableIncome": "900000", "taxPayable": {"total": "45000"}},
                                 {"taxableIncome": 0, "taxPayable": {"total": 5}},
                                 {"taxableIncome": 500000, "taxPayable": None}]):
    C("perq", LP, "calculateAdditionalTdsForPerquisite",
      loan_perquisite_helper.calculate_additional_tds_for_perquisite, p, rec)

# ================================================================== loan management internals
LMV = loan_management
for v in [None, "", " ", 0, 8.5, "9", "-1", "abc", [], True]:
    S("lm", "lm_parsePerquisiteRate", LMV.parse_perquisite_rate, v)
    S("lm", "lm_parsePerquisiteRate", LMV.parse_perquisite_rate, v, 5)
for v in [None, "", 0, 100, "2500.5", "-3", "x", True, [7]]:
    S("lm", "lm_parsePositiveAmount", LMV.parse_positive_amount, v, "loanAmount")
DATE_INPUTS = ["15/04/2025", "1/1/2025", "31/02/2025", "2025-04-15", "2025-04-15T10:00:00.000Z",
               "04/15/2025", "April 15, 2025", "garbage", "", None, " 05/06/2024 ", "29/02/2024"]
for v in DATE_INPUTS:
    S("lm", "lm_parseDateInput", LMV.parse_date_input, v, "disbursementDate")
    S("lm", "lm_formatDisplayDate", LMV.format_display_date, v)
    S("lm", "lm_buildEmiStartPeriod", LMV.build_emi_start_period, v)
EMI_BODIES = [{"emiStartMonth": "Apr", "emiStartYear": 2025}, {"emiStartMonth": "april", "emiStartYear": "2025"},
              {"emiStartMonth": 4, "emiStartYear": 2025}, {"emiStartMonth": "13", "emiStartYear": 2025},
              {"emiStartMonth": "Apr"}, {"emiStartMonth": "Apr", "emiStartYear": 1800},
              {"emiStartMonth": "Apr", "emiStartYear": "2025.5"}, {"emiStartDate": "15/04/2025"},
              {"emiStartDate": "2025-04-30T20:00:00.000Z"}, {"emiStartDate": "bad"}, {},
              {"emiStartMonth": " ", "emiStartDate": "2025-01-10"}, {"emiStartMonth": "Sept", "emiStartYear": 2025}]
for b in EMI_BODIES:
    S("lm", "lm_parseEmiStartPeriodInput", LMV.parse_emi_start_period_input, b)
for a, i in [(100000, 5000), (100000, 3000), (1000, 1000), (999.5, 100), (50000, 7000.25)]:
    S("lm", "lm_computeRepaymentSummary", LMV.compute_repayment_summary, a, i)
for d, n in itertools.product(["2025-04-01T00:00:00.000Z", "2025-11-01", None, "x"], [0, 1, 12, 20]):
    S("lm", "lm_computeLoanClosingDate", LMV.compute_loan_closing_date, d, n)
ENRICH = [dict(l, _id="L1", loanNumber="LN-00001") for l in LOANS[:4]] + [
    {"_id": "L2", "outstandingBalance": 0, "installmentAmount": 100, "closedAt": "2025-06-01T00:00:00.000Z",
     "emiStartDate": "2025-01-01T00:00:00.000Z", "status": "closed"}]
for loan, n in itertools.product(ENRICH, [0, 2]):
    S("lm", "lm_enrichLoanRecord", LMV.enrich_loan_record, loan, n)
for n, p in [(12, {"label": "Apr 2025"}), (3, {"emiStartMonthLabel": "May 2025"}), (1, {}), (2, None)]:
    S("lm", "lm_buildRepaymentMessage", LMV.build_repayment_message, n, p)
for e in [{"firstName": "Ramesh", "lastName": "Kumar"}, {"firstName": "", "payrollInfo": {"employeeCode": "E1"}},
          {"uniqueId": "BSS1"}, {}, {"firstName": " A ", "lastName": None}]:
    S("lm", "lm_resolveEmployeeDisplayName", LMV.resolve_employee_display_name, e)
for v in ["cash", "Bank Transfer", "bank-transfer", "BANK", "banktransfer", "cheque", "upi", "", None, " CASH "]:
    S("lm", "lm_normalizePaymentMode", LMV.normalize_payment_mode, v)

# ================================================================== VPF
VP = vpf_management
PREFS = [{}, {"vpf": {"enabled": True, "contributionType": "fixed", "fixedAmount": 1500, "status": "active"}},
         {"vpf": {"status": "pending", "percentage": 12, "effectiveFrom": "2025-05-01"}},
         {"vpf": {"status": "stopped", "enabled": True, "percentage": 10}},
         {"vpfEnabled": True, "vpfStatus": "active", "vpfPercentage": 15},
         {"vpfStatus": "pending", "vpfContributionType": "fixed", "vpfFixedAmount": "2000"},
         {"vpfEnabled": False}, {"vpf": "x"}, {"vpf": None, "vpfStatus": ""}]
for p in PREFS:
    C("vpf", VPF, "normalizeVpf", VP.normalize_vpf, p)
    n = VP.normalize_vpf(p)
    C("vpf", VPF, "sanitizeVpfForStatus", VP.sanitize_vpf_for_status, n)
    C("vpf", VPF, "formatVpfResponse", VP.format_vpf_response, n)
    S("vpf", "vpf_mergePayrollPreferencesWithVpf", VP.merge_payroll_preferences_with_vpf, p, n)
    S("vpf", "vpf_toVpfAdminListItem", VP.to_vpf_admin_list_item,
      {"_id": "E1", "firstName": "A", "lastName": "B", "email": "a@b", "uniqueId": "U1",
       "payrollPreferences": p, "payrollInfo": {"baseSalary": 23456}})
    for basic in [0, 15000, 23456.78, "20000", None]:
        C("vpf", VPF, "calculateMonthlyVpfAmount", VP.calculate_monthly_vpf_amount, n, basic)
VPF_EFF = [{"status": "active", "enabled": True}, {"status": "active", "effectiveFrom": "2025-05-01T00:00:00.000Z"},
           {"status": "active", "effectiveFrom": "2025-04-30T18:30:00.000Z"},
           {"status": "active", "effectiveFrom": "2025-05-31T18:29:59.999Z"},
           {"status": "pending", "effectiveFrom": "2025-01-01"}, {"status": "stopped", "enabled": True}]
for v, m, y in itertools.product(VPF_EFF, ["Apr", "May", "Jun", "Foo"], [2025, "2025"]):
    S("vpf", "vpf_isVpfEffective", VP.is_vpf_effective_for_payroll, v, m, y)
for ct, val, ack in itertools.product(["percentage", "fixed", "x", None], [0, 1, 12, 100, 101, "50", "abc"],
                                      [True, False, "yes"]):
    S("vpf", "vpf_validateStartPayload",
      lambda o: VP.validate_start_payload(o.get("contributionType", UNDEFINED),
                                          o.get("contributionValue", UNDEFINED),
                                          o.get("acknowledged", UNDEFINED)),
      {"contributionType": ct, "contributionValue": val, "acknowledged": ack})

# ================================================================== arrears / revision / import small helpers
for m in ["1", "12", "13", 0, "jan", "JAN", "January", "sept", "Sep", " Mar ", "", None, 4.0, "04"]:
    S("misc", "ar_normalizeMonthValue", arrears_management.normalize_month_value, m)
    S("misc", "pi_normalizeMonthInput", payroll_import.normalize_month_input, m)
for y in [2025, "2025", "1899", 1900, "20.5", None, "", "abc"]:
    S("misc", "pi_normalizeYearInput", payroll_import.normalize_year_input, y)
for a in [{}, {"employeeId": "E1", "approverName": "X", "approverRole": "HR", "status": "approved",
               "comments": "ok", "approvedDate": "2025-01-01T00:00:00.000Z", "_id": "z"}, None]:
    S("misc", "sr_mapApprovalForResponse", salary_revision.map_approval_for_response, a)

# ================================================================== parse import file


def _xlsx(rows, sheet_title="Sheet1", extra_sheet=False, start_cell=None):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            if val is None:
                continue
            cell = ws.cell(row=r + (start_cell or 0), column=c + (start_cell or 0), value=val)
            if isinstance(val, (datetime, date)):
                cell.number_format = "dd-mm-yyyy"
    if extra_sheet:
        wb.create_sheet("Other").cell(row=1, column=1, value="ignored")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ONE_TIME_ROWS = [
    ["Employee Code", "Category", "Component Code", "Component Name", "Amount", "Remarks"],
    ["BSSW10000", "earning", "BONUS", "Diwali Bonus", 5000, "festival"],
    ["bssw10001", "Deductions", "", "Canteen", "250.50", None],
    [None, None, None, None, None, None],
    ["E3", "benefit", "MED", None, -10, "neg"],
    ["E4", "earnings", None, None, "abc", None],
    [" E5 ", "EARNINGS", "x", "y", 0, " r "],
]
REVISION_ROWS = [
    ["employee_code", "revision_type", "effective_from", "gross_salary", "ctc", "Basic Salary",
     "HRA", "Conveyance", "Special Allowance", "Diwali Bonus", "epf", "esic", "pf_on_actuals",
     "auto_approve", "auto_implement", "arrears_applicable", "arrears_from", "state", "tax_regime",
     "payment_frequency", "remarks", None, "HRA"],
    ["BSSW10000", "Increment", datetime(2025, 4, 1), 30000, 420000, 15000, 6000, None, 100, 500,
     True, "no", "yes", "TRUE", 0, "yes", "2025-01-01", "Maharashtra", "old", "biweekly", "hike", "x", 7],
    ["E2", "hike", "45748", "25,000", "₹ 3,00,000", "", "", "", "", "", "y", 1, False, "",
     "", "", "", "", "NEW", "Monthly", None, None, None],
    ["E3", "annual", "2025-13-45", "abc", None, None, None, None, None, None, None, None, None,
     None, None, None, None, None, None, None, None, None, None],
    [None] * 23,
    ["E4", "promotion", 45748.5, 0, 0, 0, None, None, None, None, "maybe", "", None, None, None, None,
     None, None, None, None, None, None, None],
]
CSV_SAMPLES = [
    b"employee_code,category,component_code,component_name,amount,remarks\nE1,earnings,BONUS,Bonus,1000,ok\n"
    b"E2,deduction,,Canteen,\"1,200\",\"multi\nline\"\n\nE3,benefits,MED,,50,\r\n",
    b"\xef\xbb\xbfEmployee Code,Category,Amount,Amount\r\nE1,Earnings,10,20\r\nE2,x,\"he said \"\"hi\"\"\",5,extra\r\n",
    b"a,b\n1,2",
    b"",
    b"employee_code,revision_type,effective_from,gross_salary,ctc,basic,HRA,epf,esic,Unknown Col\n"
    b"E1,increment,2025-04-01,30000,,15000,,true,no,\nE2,promotion,01/04/2025,0,abc,x,5000,1,0,7\n",
]
FILES = []
for rows in (ONE_TIME_ROWS, REVISION_ROWS):
    FILES.append({"b64": base64.b64encode(_xlsx(rows)).decode(), "originalname": "import.xlsx",
                  "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
FILES.append({"b64": base64.b64encode(_xlsx(ONE_TIME_ROWS, extra_sheet=True, start_cell=1)).decode(),
              "originalname": "x.xlsx", "mimetype": "application/octet-stream"})
FILES.append({"b64": base64.b64encode(_xlsx([[2024, 1.5, True, None, "Name", "Name", "Name"],
                                              [1, 2, 3, 4, 5, 6, 7], [None, None, None, None, None, None, ""],
                                              ["=1+1", None, None, None, None, None, None]])).decode(),
              "originalname": "h.xlsx", "mimetype": "application/vnd.ms-excel"})
for data in CSV_SAMPLES:
    FILES.append({"b64": base64.b64encode(data).decode(), "originalname": "import.csv",
                  "mimetype": "text/csv"})
FILES.append({"b64": base64.b64encode(b"x").decode(), "originalname": "a.pdf", "mimetype": "application/pdf"})


def _py_file(f):
    return {"buffer": base64.b64decode(f["b64"]), "mimetype": f["mimetype"],
            "originalname": f["originalname"]}


def _py_parse(f):
    return pf.parse_payroll_import_file(_py_file(f))


for f in FILES:
    S("parse", "parseFile", _py_parse, f)
    S("parse", "parseAndMapOneTime",
      lambda f: [pf.map_import_row(r, i + 1) for i, r in enumerate(_py_parse(f))], f)
    S("parse", "parseAndMapRevision",
      lambda f: [pf.map_salary_revision_row(r, i + 1) for i, r in enumerate(_py_parse(f))], f)
ROW_SAMPLES = [{"Employee Code": "E1", "  Amount ": " 12 ", "Gross\tSalary": "1"},
               {"basic_salary": "", "Base Salary": "₹ 12,000", "new basic": 5},
               {"BASIC": "abc", "basic pay": True, "diwali_bonus": "", "hra": "1e3", "x": False},
               {"basic": 0}, {}, {"pf": 1, "esi": "Y", "lwf": 2, "pt": "nope", "pf_actuals": "0"}]
for r in ROW_SAMPLES:
    S("parse", "pf_normalizeImportRowKeys", pf.normalize_import_row_keys, r)
    nr = pf.normalize_import_row_keys(r)
    S("parse", "pf_pickBasicSalaryFromRow", pf.pick_basic_salary_from_row, nr)
    C("parse", PF, "extractComponentInputsFromRow", pf.extract_component_inputs_from_row, nr)
    C("parse", PF, "extractComponentAmountsFromRow", pf.extract_component_amounts_from_row, nr)
    C("parse", PF, "mapImportRow", pf.map_import_row, r, 3)
    C("parse", PF, "mapSalaryRevisionRow", pf.map_salary_revision_row, r, 4)
for v in [None, 5, 1.5, "", " ", "₹ 1,200.50", "$", "abc", "-7", True, "1e3"]:
    S("parse", "pf_parseOptionalImportNumber", pf.parse_optional_import_number, v)
for v in [None, True, False, 1, 0, 2, "", "Yes", " n ", "maybe", "TRUE", "1"]:
    S("parse", "pf_parseOptionalImportBoolean", pf.parse_optional_import_boolean, v)
for k in ["Basic Salary", "BASIC", "basic-pay", "New Basic Pay", "basics", "", None, "base"]:
    S("parse", "pf_isBasicImportHeader", pf.is_basic_import_header, k)

# ================================================================== one-time import helper
EMPLOYEES = [
    {"_id": "e1", "firstName": "Ramesh", "lastName": "Kumar", "customId": "NSS02", "uniqueId": "BSSW10000",
     "status": "Active", "payrollInfo": {"ctc": 360000, "grossSalary": 25000, "baseSalary": 12500,
                                         "componentOverrides": [{"category": "earnings", "name": "HRA",
                                                                 "overrideAmount": 5000}]},
     "statutoryInfo": {"ptState": "Maharashtra", "esiApplicable": True, "epfOverride": {"enabled": False}}},
    {"_id": "e2", "firstName": "Suresh", "lastName": None, "customId": "", "uniqueId": "BSSW10001",
     "status": "Active", "payrollInfo": {}, "statutoryInfo": {"lwfState": "Gujarat"}},
    {"_id": "e3", "firstName": "Old", "lastName": "", "customId": "E3", "uniqueId": "u3", "status": "Active"},
]
INACTIVE = [{"_id": "e9", "firstName": "Gone", "customId": "OLD1", "uniqueId": "BSSW19999",
             "status": "Inactive"}]
COMPS = COMPONENT_SETS["pct_gross"]
OT_ROWS = [
    {"rowNumber": 1, "employeeCode": "nss02", "category": "earnings", "componentCode": "", "componentName": "hra",
     "amount": 1000, "remarks": ""},
    {"rowNumber": 2, "employeeCode": "BSSW10001", "category": "deductions", "componentCode": "x",
     "componentName": "", "amount": 50},
    {"rowNumber": 3, "employeeCode": "OLD1", "category": "earnings", "componentName": "HRA", "amount": 1},
    {"rowNumber": 4, "employeeCode": "", "category": "earnings", "amount": 1},
    {"rowNumber": 0, "employeeCode": "E3", "category": "foo", "amount": 1},
    {"employeeCode": "E3", "category": "benefits", "componentName": "Medical Insurance", "amount": -1},
    {"employeeCode": "E3", "category": "benefits", "amount": 5},
    {"employeeCode": "nobody", "category": "benefits", "amount": 5},
    {"employeeCode": "E3", "category": "earnings", "componentCode": "conv", "amount": 0},
]
VALIDATED = one_time_import_helper.validate_one_time_import_rows(OT_ROWS, EMPLOYEES, COMPS, INACTIVE)
C("onetime", OT, "validateOneTimeImportRows", one_time_import_helper.validate_one_time_import_rows,
  OT_ROWS, EMPLOYEES, COMPS, INACTIVE)
C("onetime", OT, "validateOneTimeImportRows", one_time_import_helper.validate_one_time_import_rows,
  OT_ROWS, [], [], [])
C("onetime", OT, "summarizeRows", one_time_import_helper.summarize_rows, VALIDATED)
for ov in OVERRIDE_SETS[:5]:
    C("onetime", OT, "mergeOneTimeIntoOverrides", one_time_import_helper.merge_one_time_into_overrides,
      ov if ov is not None else [], VALIDATED)
C("onetime", OT, "applyAdHocOneTimeLines", one_time_import_helper.apply_ad_hoc_one_time_lines,
  VALIDATED + [{"valid": True, "isAdHoc": True, "category": "earnings", "amount": "99.5"},
               {"valid": True, "isAdHoc": True, "category": "deductions", "componentCode": "c", "amount": 0.4}],
  [], [], [], mutated=[1, 2, 3])
S("onetime", "groupOneTime", one_time_import_helper.group_one_time_rows_by_employee, VALIDATED)
for code in ["NSS02", "bssw10000", " e3 ", "", "OLD1", "zzz", None]:
    C("onetime", OT, "findEmployeeForImport", one_time_import_helper.find_employee_for_import,
      EMPLOYEES, code)
    C("onetime", OT, "resolveEmployeeImportError", one_time_import_helper.resolve_employee_import_error,
      code, EMPLOYEES, INACTIVE)
for cat, code, name in [("earnings", "HRA", ""), ("earnings", "", "conveyance"), ("deductions", "", "Club"),
                        ("benefits", "x", "y"), ("foo", "HRA", "HRA")]:
    C("onetime", OT, "findOrgComponent", one_time_import_helper.find_org_component, COMPS, cat, code, name)

# ================================================================== salary structure / revision import
COMPLIANCE = {"epf": {"isEnabled": True, "contributionPreferences": {"edliContribution": True,
                                                                    "adminCharges": True}},
              "esi": {"isEnabled": True}, "professionalTax": PT_BASE, "labourWelfareFund": LWF_BASE,
              "statutoryBonus": BONUS_BASE}
FIXED_DATE = "2025-04-01T00:00:00.000Z"
for comps, base, gross, ov, emp, sc in itertools.product(
        ["pct_gross", "fixed", "only_basic", "empty"], [0, 15000], [0, 30000, 12000.4],
        [None, OVERRIDE_SETS[1], OVERRIDE_SETS[3]], [None, EMPLOYEES[0]], [None, COMPLIANCE]):
    O("ssi", SSI, "computePayrollInfoFromSalaryInput", ssi.compute_payroll_info_from_salary_input,
      {"baseSalary": base, "grossSalary": gross, "salaryComponents": COMPONENT_SETS[comps],
       "overrides": ov, "employee": emp, "statutoryCompliance": sc, "state": "Maharashtra",
       "salaryEffectiveFrom": {"$date": FIXED_DATE}, "grossSalaryInput": 30001, "ctcInput": 400000,
       "taxRegime": "Old", "paymentFrequency": "Daily"})
O("ssi", SSI, "computePayrollInfoFromBaseSalary", ssi.compute_payroll_info_from_base_salary,
  {"baseSalary": 15000, "salaryComponents": COMPONENT_SETS["pct_gross"], "employee": EMPLOYEES[1],
   "statutoryCompliance": COMPLIANCE, "salaryEffectiveFrom": {"$date": FIXED_DATE}})
S("ssi", "ssi_buildRevisionSalaryComponents", ssi.build_revision_salary_components, 12000,
  [{"name": "Basic", "calculatedAmount": 12000.5}, {"name": "Special", "calculatedAmount": 500,
                                                    "isBalancing": True}, {"name": "Z"}])
S("ssi", "ssi_buildRevisionSalaryComponents", ssi.build_revision_salary_components, 9999.5, [])

def _with_emp(fn):
    def wrapper(emp, *rest):
        ret = fn(emp, *rest)
        return {"ret": ret, "emp": emp}
    return wrapper


for v in ["Diwali Bonus", " HRA-1 ", None, "", 5]:
    S("sri", "sri_normalizeExcelMatchKey", sri.normalize_excel_match_key, v)
for c in [{"code": "SPECIAL"}, {"name": "Other Allowance"}, {"code": "hra"}, {}]:
    S("sri", "sri_isBalancingAllowanceComponent", sri.is_balancing_allowance_component, c)
for cs in COMPONENT_SETS.values():
    S("sri", "sri_buildOrgComponentCatalog", sri.build_org_component_catalog, cs)
AMOUNTS = [{}, {"hra": 6000, "conveyance": 1000, "unknown_col": 5, "basic_salary": 1},
           {"HRA": "7000", "statutory bonus": "x", "annual_bonus": 0}]
BLANKS = [[], ["hra", "canteen", "foo", "medical_insurance"]]
for a, b, cs in itertools.product(AMOUNTS, BLANKS, ["pct_gross", "fixed"]):
    C("sri", SRI, "resolveExcelComponentOverrides", sri.resolve_excel_component_overrides,
      a, COMPONENT_SETS[cs], b)
for ov, cl in itertools.product(OVERRIDE_SETS[:6], [[], [{"category": "earnings", "name": "hra"}],
                                                    [{"category": "deductions", "code": "canteen"}]]):
    C("sri", SRI, "removeClearedComponentOverrides", sri.remove_cleared_component_overrides,
      ov if ov is not None else [], cl)
CAT = sri.build_org_component_catalog(COMPONENT_SETS["pct_gross"])
for bases in [{"baseSalary": 15000, "grossSalary": 30000, "ctc": 400000}, {}]:
    C("sri", SRI, "materializeDefaultAmountOverrides", sri.materialize_default_amount_overrides, CAT, bases)
S("sri", "sri_toDefaultComponentsArray", sri.to_default_components_array, OVERRIDE_SETS[1])
for f in [None, {}, {"epfApplicable": "yes", "esiApplicable": False, "pfOnActuals": "1", "ptApplicable": "x",
                     "lwfApplicable": ""}, {"esiApplicable": "no"}, "str"]:
    C("sri", SRI, "sanitizeStatutoryFlags", sri.sanitize_statutory_flags, f)
    S("sri", "applyFlags", _with_emp(sri.apply_statutory_flags_to_employee), dict(EMPLOYEES[0]), f)
    S("sri", "sri_mergeEmployeeStatutoryForCalc", sri.merge_employee_statutory_for_calc, EMPLOYEES[0], f)
for st in ["Gujarat", " ", None, "Karnataka "]:
    S("sri", "applyState", _with_emp(sri.apply_state_to_employee), dict(EMPLOYEES[1]), st)
for v in ["2025-04-01", "45748", 45748.25, "01/04/2025", "garbage", "", None]:
    S("sri", "sri_parseEffectiveDate", sri.parse_effective_date, v)
    S("sri", "sri_parseOptionalNumber", sri.parse_optional_number, v)
    S("sri", "sri_parseBoolean", sri.parse_boolean, v, True)
for v in ["Hike", "annual hike", "NEW", "promotion", "", None, "weird type"]:
    S("sri", "sri_normalizeRevisionType", sri.normalize_revision_type, v)
    C("sri", SRI, "formatRevisionTypeLabel", sri.format_revision_type_label, sri.normalize_revision_type(v))
for app, frm, eff, inc in itertools.product([True, False], ["2025-01-01", "2024-12-31T18:30:00.000Z", None],
                                            ["2025-04-01", None], [12000, 0, -5]):
    S("sri", "sri_calculateArrears", lambda o: sri.calculate_arrears(**{_snake(k): v for k, v in o.items()}),
      {"arrearsApplicable": app, "arrearsFromDate": frm, "effectiveFrom": eff, "incrementAmount": inc})
REV_RAW = [
    {"rowNumber": 1, "employeeCode": "NSS02", "revisionType": "Increment", "effectiveFrom": "2025-04-01",
     "grossSalary": 30000, "baseSalary": None, "ctcInput": 420000, "state": "", "taxRegime": "Old",
     "paymentFrequency": "", "arrearsApplicable": "true", "arrearsFromDate": "2025-01-01",
     "autoApprove": "true", "autoImplement": "", "remarks": "r",
     "componentAmounts": {"hra": 6000, "foo": 5}, "blankComponentHeaders": ["conveyance"],
     "statutoryFlags": {"esiApplicable": False, "pfOnActuals": True}},
    {"rowNumber": 2, "employeeCode": "BSSW10001", "revisionType": "promotion", "effectiveFrom": "45748",
     "grossSalary": 25000, "baseSalary": 15000, "ctcInput": None, "componentAmounts": {},
     "blankComponentHeaders": [], "statutoryFlags": {}},
    {"rowNumber": 3, "employeeCode": "E3", "revisionType": "x", "effectiveFrom": "2025-04-01",
     "grossSalary": 1, "ctcInput": None},
    {"rowNumber": 4, "employeeCode": "E3", "revisionType": "adjust", "effectiveFrom": "bad",
     "grossSalary": 1, "ctcInput": None},
    {"rowNumber": 5, "employeeCode": "E3", "revisionType": "adjust", "effectiveFrom": "2025-04-01",
     "grossSalary": "abc", "ctcInput": None},
    {"rowNumber": 6, "employeeCode": "E3", "revisionType": "adjust", "effectiveFrom": "2025-04-01",
     "grossSalary": 20000, "ctcInput": -1},
    {"rowNumber": 7, "employeeCode": "E3", "revisionType": "adjust", "effectiveFrom": "2025-04-01",
     "grossSalary": 20000, "ctcInput": None, "arrearsApplicable": "yes", "arrearsFromDate": "2025-06-01"},
    {"rowNumber": 8, "employeeCode": "E3", "revisionType": "adjust", "effectiveFrom": "2025-04-01",
     "grossSalary": 20000, "ctcInput": None, "autoImplement": "1"},
    {"rowNumber": 9, "employeeCode": "OLD1", "revisionType": "adjust"},
    {"rowNumber": 10, "employeeCode": "", "revisionType": "adjust"},
    {"rowNumber": 11, "employeeCode": "E3", "revisionType": "new joiner", "effectiveFrom": "2025-04-01",
     "grossSalary": 9000, "baseSalary": 9500, "ctcInput": 1, "arrearsApplicable": "no",
     "componentAmounts": {"basic_salary": 8000}, "statutoryFlags": {"epfApplicable": "no"}},
]
for cs, sc in itertools.product(["pct_gross", "fixed", "empty"], [None, COMPLIANCE]):
    C("sri", SRI, "validateSalaryRevisionRows", sri.validate_salary_revision_rows,
      REV_RAW, EMPLOYEES, COMPONENT_SETS[cs], sc, INACTIVE)
C("sri", SRI, "summarizeRows", sri.summarize_rows, [{"valid": True}, {"valid": False}, {}])


@unittest.skipUnless(node_available(), SKIP_REASON)
class LoansRevisionsImportsParityTests(SimpleTestCase):
    def _check(self, group):
        cases = CASES[group]
        js_results = run_js_batch([c.js_call() for c in cases], timeout=300)
        self.assertEqual(len(js_results), len(cases))
        failures = []
        for case, js in zip(cases, js_results):
            py = run_py(case)
            if is_error(js) or is_error(py):
                if not (is_error(js) and is_error(py)):
                    failures.append(f"{case.label()}\n    JS: {txt(js)[:400]}\n    PY: {txt(py)[:400]}")
                continue
            if txt(js) != txt(py):
                a, b = txt(js), txt(py)
                i = next((k for k in range(min(len(a), len(b))) if a[k] != b[k]), min(len(a), len(b)))
                failures.append(f"{case.label()}\n    JS: ...{a[max(0, i - 120):i + 200]}\n"
                                f"    PY: ...{b[max(0, i - 120):i + 200]}")
        self.assertFalse(failures, f"{len(failures)}/{len(cases)} parity failures in {group}:\n"
                         + "\n".join(failures[:15]))

    def test_loan_helper(self):
        self._check("loan")

    def test_loan_perquisite_helper(self):
        self._check("perq")

    def test_loan_management_helpers(self):
        self._check("lm")

    def test_vpf_helpers(self):
        self._check("vpf")

    def test_month_year_and_approval_helpers(self):
        self._check("misc")

    def test_parse_payroll_import_file(self):
        self._check("parse")

    def test_one_time_import_helper(self):
        self._check("onetime")

    def test_salary_structure_import_helper(self):
        self._check("ssi")

    def test_salary_revision_import_helper(self):
        self._check("sri")
