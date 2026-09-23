"""Employer → contractor SAP populate endpoints.

* ``GET  /v1/employer/sap/config``  — whether SAP is configured + the field map.
* ``POST /v1/employer/contractors/<id>/sap-sync`` — pull the contractor's master
  data from SAP (live OData) or from an uploaded SAP export (``sap`` file) and
  map it onto the contractor profile.
"""

# ruff: noqa: F401,F403
import re

from .common import *


def _profile_payload(profile):
    return {
        "contractorCode": profile.contractor_code,
        "contractorName": profile.contractor_name,
        "tradeName": profile.trade_name,
        "gstin": profile.gstin,
        "panNo": profile.pan_no,
        "cin": profile.cin,
        "mobile": profile.mobile,
        "email": profile.email,
        "address": profile.address,
        "district": profile.district,
        "state": profile.state,
        "clraLicenceNo": profile.clra_licence_no,
        "registrationSource": profile.registration_source,
    }


def _settings_payload(cfg, employer=None):
    from api import sap_client
    return {
        "configured": bool(cfg.get("baseUrl")),
        "mode": cfg.get("mode") or "b1",
        "authType": cfg.get("authType") or "session",
        "baseUrl": cfg.get("baseUrl") or "",
        "companyDb": cfg.get("companyDb") or "",
        "username": cfg.get("username") or "",
        "entity": cfg.get("entity") or "",
        "idField": cfg.get("idField") or "",
        "bpFilter": cfg.get("bpFilter") or "",
        "verifyTls": bool(cfg.get("verifyTls", True)),
        "extraLogin": cfg.get("extraLogin") or {},
        "fieldMapping": cfg.get("fieldMapping") or {},
        "hasPassword": bool(cfg.get("password")),
        "hasToken": bool(cfg.get("token")),
        "allowlistIps": sap_client.allowlist_ips(cfg),
        "fieldMap": sap_client.field_map_payload(),
    }


@api_view(["GET", "PUT"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_settings(request):
    """Employer SAP connection settings (secrets encrypted; masked on read)."""
    from api import sap_client
    from api import crypto
    from bss_core.models import EmployerSapConfig
    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)

    row = EmployerSapConfig.objects.filter(employer=user).first()
    if request.method == "PUT":
        body = request.data if isinstance(request.data, dict) else {}
        if row is None:
            row = EmployerSapConfig(employer=user)
        row.mode = "b1" if str(body.get("mode") or row.mode) == "b1" else "odata"
        row.auth_type = str(body.get("authType") or row.auth_type or "session")
        if row.auth_type not in ("session", "bearer", "basic"):
            row.auth_type = "session"
        row.base_url = str(body.get("baseUrl") or "")[:255]
        row.company_db = str(body.get("companyDb") or "")[:80]
        row.username = str(body.get("username") or "")[:120]
        row.entity = str(body.get("entity") or "A_BusinessPartner")[:80]
        row.id_field = str(body.get("idField") or "CardCode")[:80]
        row.bp_filter = str(body.get("bpFilter") or "")[:255]
        row.verify_tls = bool(body.get("verifyTls", True))
        row.allowlist_ips = str(body.get("allowlistIps") or "")[:255]
        extra = body.get("extraLogin")
        row.extra_login = extra if isinstance(extra, dict) else {}
        fm = body.get("fieldMapping")
        row.field_mapping = fm if isinstance(fm, dict) else {}
        # Secrets: only overwrite when a non-empty new value is supplied.
        if str(body.get("password") or "").strip():
            row.password_enc = crypto.encrypt(str(body["password"]))
        if str(body.get("token") or "").strip():
            row.token_enc = crypto.encrypt(str(body["token"]))
        if body.get("clearSecrets") is True:
            row.password_enc = ""
            row.token_enc = ""
        row.is_active = bool(body.get("isActive", True))
        row.save()
        return Response({"data": _settings_payload(sap_client.resolve_config(user))},
                        status=status.HTTP_200_OK)

    return Response({"data": _settings_payload(sap_client.resolve_config(user), user)},
                    status=status.HTTP_200_OK)


@api_view(["POST"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_test(request):
    """Test the employer's SAP connection: log in + fetch one Business Partner."""
    from api import sap_client
    from bss_core.models import EmployerSapConfig
    from django.utils import timezone
    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)
    cfg = sap_client.resolve_config(user)
    if not cfg.get("baseUrl"):
        return Response({"error": "Save the SAP connection URL first."},
                        status=status.HTTP_400_BAD_REQUEST)
    identifier = str((request.data or {}).get("identifier") or "").strip()
    if not identifier:
        return Response({"error": "Provide a Contractor code / CIN / Vendor no. to test."},
                        status=status.HTTP_400_BAD_REQUEST)
    row = EmployerSapConfig.objects.filter(employer=user).first()
    try:
        record = sap_client.fetch_record(identifier, cfg=cfg)
        mapped = sap_client.map_sap_record(record, cfg.get("fieldMapping"))
        if row is not None:
            row.last_tested_at = timezone.now()
            row.last_status = f"OK · {len(mapped)} field(s)"
            row.save(update_fields=["last_tested_at", "last_status"])
        return Response({"data": {"ok": True, "mapped": sorted(mapped.keys()),
                                  "mappedCount": len(mapped), "source": cfg.get("mode")}},
                        status=status.HTTP_200_OK)
    except RuntimeError as exc:
        if row is not None:
            row.last_tested_at = timezone.now()
            row.last_status = f"FAILED · {str(exc)[:180]}"
            row.save(update_fields=["last_tested_at", "last_status"])
        return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["GET"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_config(request):
    """SAP integration status + the SAP→profile field map (for the UI)."""
    from api import sap_client
    if request.user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)
    return Response({"data": _settings_payload(sap_client.resolve_config(request.user))},
                    status=status.HTTP_200_OK)


def _setup_document_html(employer_name, cfg, ips, field_map):
    """Printable setup + IP-whitelisting guide for one employer."""
    from html import escape
    rows = "".join(
        f"<tr><td>{escape(m['label'])}</td><td><code>{escape(m['field'])}</code></td>"
        f"<td>{escape(', '.join(m['sapFields']))}</td></tr>" for m in field_map)
    ip_list = "".join(f"<li><code>{escape(ip)}</code></li>" for ip in ips) or "<li><em>not detected</em></li>"
    mode = cfg["mode"]
    sap_hint = ("SAP Business One <strong>Service Layer</strong> (usually port <code>50000</code>, "
                "e.g. <code>https://sap-host:50000/b1s/v1</code>)"
                if mode == "b1" else
                "SAP Gateway <strong>OData</strong> service (e.g. <code>/sap/opu/odata/sap/API_BUSINESS_PARTNER</code>)")
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>BSS SAP Setup — {escape(employer_name)}</title>
<style>
 body {{ font-family: Arial, Helvetica, sans-serif; margin: 32px; color:#111; line-height:1.5; }}
 h1 {{ font-size:20px; margin:0 0 4px; }} h2 {{ font-size:14px; margin:22px 0 6px; text-transform:uppercase; letter-spacing:.04em; }}
 .sub {{ color:#666; font-size:12px; margin-bottom:16px; }}
 code, pre {{ font-family: ui-monospace, Menlo, Consolas, monospace; }}
 pre {{ background:#f4f6f8; border:1px solid #d7dde3; border-radius:6px; padding:10px; font-size:12px; white-space:pre-wrap; }}
 table {{ width:100%; border-collapse:collapse; font-size:12px; }} th,td {{ border:1px solid #cfd6dd; padding:6px 8px; text-align:left; vertical-align:top; }}
 th {{ background:#0A192F; color:#fff; }}
 ul {{ margin:6px 0 0 18px; }}
 .tag {{ display:inline-block; font-size:11px; font-weight:700; padding:2px 6px; border-radius:4px; background:#e8f0fb; color:#0b4a8b; }}
 .warn {{ background:#fff7ed; border:1px solid #fed7aa; color:#9a3412; padding:8px 10px; border-radius:6px; font-size:12px; }}
</style></head><body>
 <h1>SAP Integration — Setup &amp; IP Whitelisting</h1>
 <div class="sub">Prepared for <strong>{escape(employer_name)}</strong> · Generated {escape(str(__import__('datetime').date.today()))} · Bharat Shram Setu</div>

 <h2>1. Allow-list the BSS server IP(s) at your SAP</h2>
 <p>Add the following outbound IP{'' if len(ips)==1 else 's'} to the firewall / reverse proxy / SAP access rules that protect your SAP API, so the BSS platform can call it:</p>
 <ul>{ip_list}</ul>
 <div class="warn">If your SAP sits behind a VPN/Tailscale or an SAP Cloud Connector, enable that path for these IPs instead of a public firewall rule. Allow the SAP port (B1 Service Layer default <code>50000</code>; Gateway <code>443</code>).</div>

 <h2>2. Provide a SAP API user</h2>
 <p>Create a technical/communication user with read access to Business Partners (B1) / the OData service, and share:</p>
 <pre>CompanyDB (if SAP B1): {escape(cfg['companyDb'] or '&lt;company database&gt;')}
Username: {escape(cfg['username'] or '&lt;api user&gt;')}
Password: (shared separately, never over email)
Service base URL: {escape(cfg['baseUrl'] or '&lt;https://sap-host:50000/b1s/v1&gt;')}</pre>

 <h2>3. Enable the API service</h2>
 <p>Target: {sap_hint}. Mode configured: <span class="tag">{escape(mode)}</span>.</p>
 <p>BSS will log in and read a Business Partner by its identifier (CardCode / Vendor number / BusinessPartner). Typical call:</p>
 <pre>POST {escape(cfg['baseUrl'] or '&lt;base&gt;')}/{ 'Login' if mode=='b1' else cfg['entity'] }
GET  {escape(cfg['baseUrl'] or '&lt;base&gt;')}/{ "BusinessPartners('VEND01519')" if mode=='b1' else (cfg['entity'] + "?$filter=" + cfg['idField'] + " eq '<id>'") }</pre>

 <h2>4. Field mapping (SAP → BSS contractor profile)</h2>
 <table><thead><tr><th>BSS field</th><th>Key</th><th>SAP field(s) accepted</th></tr></thead><tbody>{rows}</tbody></table>

 <h2>5. How to populate</h2>
 <p>In BSS: <strong>Employer → My Contractors → Populate from SAP</strong> → enter the contractor&rsquo;s CardCode/Vendor number (or upload a SAP export CSV/XLSX). The mapped fields are written into the contractor profile.</p>
</body></html>"""


@api_view(["GET"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_setup_document(request):
    """Download the SAP setup & IP-whitelisting document for this employer."""
    from django.http import HttpResponse
    from api import sap_client
    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)
    emp_profile = getattr(user, "employer_profile", None)
    employer_name = (emp_profile.employer_name if emp_profile else "") or str(user)
    cfg = sap_client.resolve_config(user)
    html = _setup_document_html(employer_name, cfg,
                                sap_client.allowlist_ips(cfg), sap_client.field_map_payload())
    filename = f"bss-sap-setup-{re.sub(r'[^A-Za-z0-9]+', '-', employer_name).strip('-') or 'employer'}.html"
    response = HttpResponse(html, content_type="text/html; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@api_view(["POST"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_contractor_sap_sync(request, contractor_id):
    """Populate a bound contractor's profile from SAP (live or uploaded export)."""
    from django.contrib.auth import get_user_model
    from api import sap_client
    from bss_core.models import ContractorEmployerBinding, ContractorProfile
    from .admin_console import log_audit

    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)

    User = get_user_model()
    contractor = User.objects.filter(id=contractor_id, role="contractor").first()
    if contractor is None:
        return Response({"error": "Contractor not found."}, status=status.HTTP_404_NOT_FOUND)

    bound = (ContractorEmployerBinding.objects
             .filter(employer=user, contractor=contractor)
             .exclude(binding_status__in=["REJECTED", "CANCELLED"]).exists())
    if not bound:
        return Response({"error": "This contractor is not bound to your organisation."},
                        status=status.HTTP_403_FORBIDDEN)

    profile, _ = ContractorProfile.objects.get_or_create(user=contractor)
    cfg = sap_client.resolve_config(user)
    identifier = str(request.data.get("identifier") or profile.contractor_code
                     or getattr(contractor, "bss_id", "") or "").strip()

    upload = request.FILES.get("sap")
    source = ""
    try:
        if upload is not None:
            record = sap_client.parse_sap_export(upload, identifier)
            source = f"file:{os.path.basename(getattr(upload, 'name', '') or 'sap-export')}"
        else:
            if not identifier:
                return Response({"error": "No Contractor code / CIN to look up in SAP."},
                                status=status.HTTP_400_BAD_REQUEST)
            if not cfg.get("baseUrl"):
                return Response({"error": "SAP is not configured for your organisation yet — "
                                          "set it up in SAP Settings."},
                                status=status.HTTP_400_BAD_REQUEST)
            record = sap_client.fetch_record(identifier, cfg=cfg)
            source = "sap-b1" if cfg.get("mode") == "b1" else "sap-odata"
    except RuntimeError as exc:
        return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    if not record:
        return Response({"error": "No SAP record found."}, status=status.HTTP_404_NOT_FOUND)

    mapped = sap_client.map_sap_record(record, cfg.get("fieldMapping"))
    if not mapped:
        return Response({"error": "The SAP record had no recognisable contractor fields."},
                        status=status.HTTP_400_BAD_REQUEST)

    for field, value in mapped.items():
        if hasattr(profile, field):
            setattr(profile, field, value)

    employer_name = ""
    emp_profile = getattr(user, "employer_profile", None)
    if emp_profile is not None:
        employer_name = emp_profile.employer_name or ""
    profile.registration_source = source[:80]
    profile.registered_by = (employer_name or str(user))[:120]
    profile.save()

    log_audit(user, "UPDATE", "contractor_profile", contractor.id,
              after={"summary": f"SAP populate for {profile.contractor_name or contractor_id} ({source})"})

    return Response({"data": {
        "profile": _profile_payload(profile),
        "mapped": sorted(mapped.keys()),
        "source": source,
    }}, status=status.HTTP_200_OK)


def _bound_contractors(user):
    """Bound contractors of this employer with BSS id, name, sites and current
    SAP mapping status."""
    from bss_core.models import (User, WorkerContractorBinding, ContractorEmployerBinding,
                                 ContractorSiteAllotment, EmployerSapContractorMap)
    cids = set(WorkerContractorBinding.objects.filter(
        site__employer=user, binding_status="ACCEPTED").values_list("contractor_id", flat=True))
    cids |= set(ContractorEmployerBinding.objects.filter(
        employer=user, binding_status="ACCEPTED").values_list("contractor_id", flat=True))
    maps = {m.contractor_id: m for m in EmployerSapContractorMap.objects.filter(employer=user)}
    out = []
    for c in User.objects.filter(id__in=cids, role="contractor"):
        cp = getattr(c, "contractor_profile", None)
        sites = set(ContractorSiteAllotment.objects.filter(
            binding__contractor=c, binding__employer=user,
            binding__binding_status="ACCEPTED", status="APPROVED").values_list("site__site_name", flat=True))
        sites |= set(WorkerContractorBinding.objects.filter(
            contractor=c, site__employer=user, binding_status="ACCEPTED").values_list("site__site_name", flat=True))
        row = maps.get(c.id)
        if row is None:
            status = "NOT_MAPPED"
        else:
            status = row.status if row.sap_id else "NOT_MAPPED"
        out.append({
            "id": c.id,
            "bssId": payloads._bss_contractor_code(c.id),
            "name": (cp.contractor_name if cp and cp.contractor_name else (c.get_full_name() or c.username)),
            "sites": sorted(s for s in sites if s),
            "sapId": (row.sap_id if row else ""),
            "status": status,
            "lastStatus": (row.last_status if row else ""),
            "lastSyncedAt": (payloads._dts(row.last_synced_at) if row and row.last_synced_at else ""),
            "mappedFields": (row.mapped_fields if row else []),
        })
    out.sort(key=lambda r: (r["name"] or "").lower())
    return out


_STATUS_LABEL = {"NOT_MAPPED": "Not mapped", "PENDING": "Pending pull",
                 "SYNCED": "Synced", "FAILED": "Failed"}


@api_view(["GET"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_bulk_template(request):
    """Download the XLSX template: bound contractors with a blank SAP ID column."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from django.http import HttpResponse

    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)

    wb = Workbook()
    ws = wb.active
    ws.title = "Contractors"
    headers = ["BSS Contractor ID", "Contractor Name", "SAP ID", "Sites", "Status"]
    fill = PatternFill("solid", fgColor="0A192F")
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        ws.column_dimensions[cell.column_letter].width = [22, 30, 24, 34, 14][col - 1]
    for r, row in enumerate(_bound_contractors(user), start=2):
        ws.cell(row=r, column=1, value=row["bssId"])
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=row["sapId"])
        ws.cell(row=r, column=4, value=", ".join(row["sites"]))
        ws.cell(row=r, column=5, value=_STATUS_LABEL.get(row["status"], row["status"]))

    info = wb.create_sheet("Instructions")
    for i, line in enumerate([
        "BSS — Bulk Contractor SAP Pull",
        "",
        "1. Fill the 'SAP ID' column for each contractor (leave the BSS ID as-is).",
        "   SAP ID = the contractor's Business Partner id: CardCode (SAP Business One)",
        "   or BusinessPartner (SAP Gateway OData).",
        "2. Save as .xlsx and upload it in Employer → SAP Settings → Bulk contractor pull.",
        "3. Choose the pull option (Skip already mapped / Force re-pull / Retry failed / Pull pending).",
        "4. Mappings are saved first; BSS then pulls each contractor's master data from SAP.",
        "",
        "Status legend: Not mapped · Pending pull · Synced · Failed.",
        "Already-mapped rows are skipped by default — switch the option to re-pull.",
    ], start=1):
        info.cell(row=i, column=1, value=line)
    info.column_dimensions["A"].width = 100

    out = io.BytesIO()
    wb.save(out)
    response = HttpResponse(out.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="bss-sap-bulk-mapping.xlsx"'
    return response


@api_view(["GET"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_bulk_status(request):
    """Live status of bound contractors: mapped / pulled / pending / failed."""
    if request.user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)
    rows = _bound_contractors(request.user)
    counts = {"mapped": 0, "synced": 0, "pending": 0, "failed": 0, "notMapped": 0}
    for r in rows:
        st = r["status"]
        if st == "SYNCED":
            counts["mapped"] += 1; counts["synced"] += 1
        elif st == "PENDING":
            counts["mapped"] += 1; counts["pending"] += 1
        elif st == "FAILED":
            counts["mapped"] += 1; counts["failed"] += 1
        else:
            counts["notMapped"] += 1
    return Response({"data": {"rows": rows, "counts": counts}}, status=status.HTTP_200_OK)


def _parse_mapping_rows(file_obj):
    """Read the filled mapping template → [{bssId, sapId}]."""
    from openpyxl import load_workbook
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb["Contractors"] if "Contractors" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c or "").strip().lower() for c in rows[0]]
    idx_bss = next((i for i, h in enumerate(header) if "bss" in h), 0)
    idx_sap = next((i for i, h in enumerate(header) if "sap" in h), 2)
    out = []
    for values in rows[1:]:
        if not values:
            continue
        bss = str(values[idx_bss] or "").strip() if idx_bss < len(values) else ""
        sap = str(values[idx_sap] or "").strip() if idx_sap < len(values) else ""
        if not bss and not sap:
            continue
        out.append({"bssId": bss, "sapId": sap})
    return out


def _contractor_id_from_bss(bss_id):
    match = re.match(r"^BSS-C0*(\d+)$", str(bss_id or "").strip(), re.I)
    return int(match.group(1)) if match else None


@api_view(["POST"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_bulk_sync(request):
    """Bulk-map BSS contractors to SAP ids and pull each from SAP."""
    from api import sap_client
    from datetime import datetime
    from django.utils import timezone
    from bss_core.models import (User, ContractorProfile, EmployerSapConfig,
                                 EmployerSapContractorMap, SapSyncLog)
    from .admin_console import log_audit

    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)
    upload = request.FILES.get("mapping")
    if upload is None:
        return Response({"error": "Upload the filled mapping template (.xlsx)."},
                        status=status.HTTP_400_BAD_REQUEST)
    mode = str(request.data.get("mode") or "skip").strip().lower()
    if mode not in ("skip", "force", "failed", "pending"):
        mode = "skip"

    try:
        rows = _parse_mapping_rows(upload)
    except Exception as exc:  # noqa: BLE001
        return Response({"error": f"Could not read the template: {exc}"},
                        status=status.HTTP_400_BAD_REQUEST)
    if not rows:
        return Response({"error": "No contractor rows found in the template."},
                        status=status.HTTP_400_BAD_REQUEST)

    bound = {r["id"]: r for r in _bound_contractors(user)}
    by_bss = {r["bssId"].upper(): r for r in bound.values()}
    cfg = sap_client.resolve_config(user)
    live = bool(cfg.get("baseUrl"))
    existing = {m.contractor_id: m for m in EmployerSapContractorMap.objects.filter(employer=user)}

    results = []
    seen = set()
    for row in rows:
        bss = row["bssId"].upper()
        sap_id = row["sapId"]
        entry = {"bssId": row["bssId"], "sapId": sap_id, "contractor": "", "status": "skipped", "detail": ""}
        info = by_bss.get(bss)
        if info is None:
            cid = _contractor_id_from_bss(row["bssId"])
            info = bound.get(cid) if cid else None
        if info is None:
            entry["detail"] = "BSS ID is not a contractor bound to your organisation."
            entry["status"] = "failed"
            results.append(entry); continue
        if bss in seen:
            entry["detail"] = "Duplicate BSS ID in the file."
            entry["status"] = "failed"
            entry["contractor"] = info["name"]
            results.append(entry); continue
        seen.add(bss)
        entry["contractor"] = info["name"]
        cid = info["id"]

        current = existing.get(cid)
        current_sap = (current.sap_id if current else "") or ""
        eff_sap = sap_id or current_sap
        if not eff_sap:
            entry["detail"] = "No SAP ID provided."
            entry["status"] = "skipped"
            results.append(entry); continue

        should_pull = True
        if not sap_id and current_sap:
            should_pull = False  # no change to mapping
        if mode == "skip" and current_sap:
            should_pull = False
        if mode == "force":
            should_pull = True
        if mode == "failed":
            should_pull = (current is None or current.status != "FAILED")
            should_pull = not should_pull  # only failed rows
        if mode == "pending":
            should_pull = (current is None or current.status == "SYNCED")
            should_pull = not should_pull  # only non-synced

        # Persist the mapping first (always).
        if current is None:
            current = EmployerSapContractorMap(employer=user, contractor_id=cid)
        if sap_id:
            current.sap_id = sap_id
        if not should_pull:
            if not current.pk:
                current.status = "PENDING"
                current.last_status = "Mapping saved — pull pending."
            current.save()
            existing[cid] = current
            entry["status"] = "skipped"
            entry["detail"] = "Already mapped (use Force re-pull to refresh)." if current_sap else "Mapping saved; not pulled by the selected option."
            results.append(entry); continue

        current.save()
        existing[cid] = current

        if not live:
            current.status = "PENDING"
            current.last_status = "SAP not configured — mapping saved, pull pending."
            current.save(update_fields=["status", "last_status", "updated_at"])
            entry["status"] = "pending"
            entry["detail"] = "Mapping saved. Configure SAP, then use Retry failed / Pull pending to sync."
            results.append(entry); continue

        try:
            record = sap_client.fetch_record(eff_sap, cfg=cfg)
            mapped = sap_client.map_sap_record(record, cfg.get("fieldMapping"))
            if not mapped:
                raise RuntimeError("SAP record had no recognisable contractor fields.")
            profile, _ = ContractorProfile.objects.get_or_create(user_id=cid)
            for field, value in mapped.items():
                if hasattr(profile, field):
                    setattr(profile, field, value)
            emp_profile = getattr(user, "employer_profile", None)
            profile.registration_source = f"sap-{cfg.get('mode')}-bulk"[:80]
            profile.registered_by = ((emp_profile.employer_name if emp_profile else "") or str(user))[:120]
            profile.save()
            current.status = "SYNCED"
            current.last_source = f"sap-{cfg.get('mode')}"
            current.last_status = f"OK · {len(mapped)} field(s)"
            current.mapped_fields = sorted(mapped.keys())
            current.last_synced_at = timezone.now()
            current.save()
            SapSyncLog.objects.create(employer=user, contractor_id=cid,
                                      source=current.last_source, status="SYNCED",
                                      mapped={k: str(v) for k, v in mapped.items()},
                                      message=f"{len(mapped)} field(s)")
            log_audit(user, "UPDATE", "contractor_profile", cid,
                      after={"summary": f"Bulk SAP pull for {info['name']} ({current.last_source})"})
            entry["status"] = "synced"
            entry["detail"] = f"{len(mapped)} field(s): " + ", ".join(sorted(mapped.keys())[:8])
        except Exception as exc:  # noqa: BLE001
            current.status = "FAILED"
            current.last_status = str(exc)[:200]
            current.save(update_fields=["status", "last_status", "updated_at"])
            SapSyncLog.objects.create(employer=user, contractor_id=cid, source=f"sap-{cfg.get('mode')}",
                                      status="FAILED", mapped={}, message=str(exc)[:255])
            entry["status"] = "failed"
            entry["detail"] = str(exc)[:180]
        results.append(entry)

    summary = {"total": len(results), "synced": 0, "skipped": 0, "failed": 0, "pending": 0}
    for r in results:
        summary[r["status"]] = summary.get(r["status"], 0) + 1
    return Response({"data": {"results": results, "summary": summary, "live": live, "mode": mode}},
                    status=status.HTTP_200_OK)


@api_view(["GET"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def employer_sap_bulk_setup_document(request):
    """Printable guide for the employer's SAP engineer to enable bulk pull."""
    from django.http import HttpResponse
    from html import escape
    from api import sap_client

    user = request.user
    if user.role != "employer":
        return Response({"error": "Employers only."}, status=status.HTTP_403_FORBIDDEN)
    cfg = sap_client.resolve_config(user)
    ips = sap_client.allowlist_ips(cfg)
    mode = cfg.get("mode") or "b1"
    ip_html = "".join(f"<li><code>{escape(ip)}</code></li>" for ip in ips) or "<li><em>not detected</em></li>"
    base = escape(cfg.get("baseUrl") or "<https://sap-host:50000/b1s/v1>")
    entity = escape(cfg.get("entity") or "BusinessPartners")
    id_field = escape(cfg.get("idField") or ("CardCode" if mode == "b1" else "BusinessPartner"))
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>BSS SAP Bulk Pull — Setup</title>
<style>body{{font-family:Arial,Helvetica,sans-serif;margin:32px;color:#111;line-height:1.5}}
h1{{font-size:20px;margin:0 0 4px}}h2{{font-size:14px;margin:22px 0 6px;text-transform:uppercase;letter-spacing:.04em}}
.sub{{color:#666;font-size:12px;margin-bottom:16px}}code,pre{{font-family:ui-monospace,Menlo,Consolas,monospace}}
pre{{background:#f4f6f8;border:1px solid #d7dde3;border-radius:6px;padding:10px;font-size:12px;white-space:pre-wrap}}
table{{width:100%;border-collapse:collapse;font-size:12px}}th,td{{border:1px solid #cfd6dd;padding:6px 8px;text-align:left}}
th{{background:#0A192F;color:#fff}}ul{{margin:6px 0 0 18px}}.warn{{background:#fff7ed;border:1px solid #fed7aa;color:#9a3412;padding:8px 10px;border-radius:6px;font-size:12px}}</style>
</head><body>
<h1>SAP Integration — Bulk Contractor Pull Setup</h1>
<div class="sub">For the employer's SAP administrator · Bharat Shram Setu</div>
<h2>1. Allow-list the BSS server IP(s)</h2><ul>{ip_html}</ul>
<div class="warn">Allow the SAP API port (B1 Service Layer <code>50000</code>; Gateway <code>443</code>). For VPN / Tailscale / SAP Cloud Connector, enable that path for the BSS IPs.</div>
<h2>2. Provide a read-only API user</h2>
<pre>CompanyDB (B1): {escape(cfg.get('companyDb') or '<company database>')}
Username:       {escape(cfg.get('username') or '<api user>')}
Password:       (shared separately)
Base URL:       {base}</pre>
<h2>3. Business Partner lookup</h2>
<p>BSS resolves each contractor by its SAP ID — <code>{id_field}</code> — using the configured {"Service Layer" if mode=='b1' else "OData"} endpoint:</p>
<pre>POST {base}/{ 'Login' if mode=='b1' else '' }
GET  {base}/{ (entity + "('" + id_field + "')") if mode=='b1' else (entity + "?$filter=" + id_field + " eq '<SAP ID>'") }</pre>
<h2>4. How the bulk template maps to SAP</h2>
<p>The employer fills the <strong>BSS ↔ SAP ID</strong> mapping template (XLSX). Each row = one bound contractor:</p>
<table><thead><tr><th>BSS Contractor ID</th><th>Contractor Name</th><th>SAP ID</th></tr></thead>
<tbody><tr><td>BSS-C00000000004</td><td>BuildWell Contractors Pvt Ltd</td><td>{escape('<CardCode / BusinessPartner>')}</td></tr></tbody></table>
<h2>5. Fields BSS reads</h2>
<p>Name, trade name, GSTIN, PAN, CIN, registration no., addresses/state, CLRA licence no., contacts… (see the full field map in SAP Settings). Unmapped custom fields can be pinned via per-employer field mapping.</p>
<h2>6. Steps</h2>
<ol><li>Allow-list the BSS IP(s).</li><li>Create the read-only user.</li><li>Share CompanyDB/username/base URL with the employer.</li>
<li>Employer enters them in SAP Settings → Test.</li>
<li>Employer fills the bulk template and uploads it → BSS pulls all.</li></ol>
</body></html>"""
    response = HttpResponse(html, content_type="text/html; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="bss-sap-bulk-setup.html"'
    return response


__all__ = ["employer_sap_config", "employer_sap_settings", "employer_sap_test",
           "employer_contractor_sap_sync", "employer_sap_setup_document",
           "employer_sap_bulk_template", "employer_sap_bulk_status",
           "employer_sap_bulk_sync", "employer_sap_bulk_setup_document"]
