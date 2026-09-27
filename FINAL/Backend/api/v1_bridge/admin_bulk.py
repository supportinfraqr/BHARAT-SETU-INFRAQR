"""BSS Admin bulk registration (password-protected).

Four XLSX-driven bulk routes — employers, contractors, workers and the
relations master table that links them (employer ↔ contractor ↔ worker + site).
Every request must carry the BSS Admin session AND the bulk-registration
password (env ``BSS_BULK_REGISTER_PASSWORD``).
"""

# ruff: noqa: F401,F403
import io
import os

from .common import *  # noqa: F401,F403


# Entity → (sheet title, [(column header, field key), ...])
ENTITIES = {
    "employers": ("Employers", [
        ("Employer Name", "name"), ("Mobile", "mobile"), ("Email", "email"),
        ("Entity Type", "entity_type"), ("CIN", "cin"), ("GSTIN", "gstin"),
        ("PAN", "pan"), ("Address", "address"), ("District", "district"),
        ("City", "city"), ("State", "state"), ("PIN", "pincode"),
        ("Contact Person", "contact_person"),
    ]),
    "contractors": ("Contractors", [
        ("Contractor Name", "name"), ("Mobile", "mobile"), ("Email", "email"),
        ("Trade Name", "trade_name"), ("Constitution", "constitution"),
        ("CIN", "cin"), ("GSTIN", "gstin"), ("PAN", "pan"),
        ("CLRA Licence No", "clra_licence_no"), ("Licence Valid To", "licence_valid_to"),
        ("Operating States", "operating_states"), ("Address", "address"),
        ("District", "district"), ("City", "city"), ("State", "state"),
    ]),
    "workers": ("Workers", [
        ("Worker Name", "name"), ("Mobile", "mobile"), ("Date of Birth", "dob"),
        ("Gender", "gender"), ("Father Name", "father_name"), ("Aadhaar", "aadhaar"),
        ("PAN", "pan"), ("UAN", "uan"), ("ESIC IP", "esic"), ("Address", "address"),
        ("District", "district"), ("City", "city"), ("State", "state"), ("PIN", "pincode"),
    ]),
    "relations": ("Relations", [
        ("Employer Mobile", "employer_mobile"), ("Contractor Mobile", "contractor_mobile"),
        ("Worker Mobile", "worker_mobile"), ("Site Name", "site_name"),
        ("Site Code", "site_code"), ("Designation", "designation"), ("Wage", "wage"),
    ]),
}


def _bulk_password():
    return os.environ.get("BSS_BULK_REGISTER_PASSWORD") or "bss-bulk-2026"


def _authorize(request):
    """Return an error Response, or None when authorised."""
    user = request.user
    if user.role != "bss_admin":
        return Response({"error": "BSS Admin only."}, status=status.HTTP_403_FORBIDDEN)
    supplied = str(request.data.get("password") or request.query_params.get("password") or "")
    if supplied != _bulk_password():
        return Response({"error": "Invalid bulk-registration password."},
                        status=status.HTTP_403_FORBIDDEN)
    return None


def _as_date(value):
    from datetime import datetime
    text = str(value or "").strip()
    if isinstance(value, datetime):
        return value.date()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _mobile(value):
    return _normalize_mobile(str(value or "").strip()) or str(value or "").strip()


def _rows(file_obj, entity):
    from openpyxl import load_workbook
    headers = [h for h, _ in ENTITIES[entity][1]]
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb[ENTITIES[entity][0]] if ENTITIES[entity][0] in wb.sheetnames else wb.active
    data = list(ws.iter_rows(values_only=True))
    if not data:
        return []
    # Map the file header row to keys (case/space-insensitive), fall back to order.
    file_header = [str(c or "").strip().lower().replace(" ", "") for c in data[0]]
    keys = [k for _, k in ENTITIES[entity][1]]
    index = {}
    for pos, (label, key) in enumerate(ENTITIES[entity][1]):
        norm = label.lower().replace(" ", "")
        index[pos] = file_header.index(norm) if norm in file_header else pos
    out = []
    for values in data[1:]:
        if not values or not any(str(v or "").strip() for v in values):
            continue
        row = {}
        for pos, key in enumerate(keys):
            i = index.get(pos, pos)
            row[key] = values[i] if i < len(values) else None
        out.append(row)
    return out


def _find_user_by_mobile(mobile):
    from bss_core.models import UserCredential
    if not mobile:
        return None
    cred = (UserCredential.objects.filter(identity_type__in=["mobile", "phone"],
                                          identity_value=mobile).select_related("user").first())
    return cred.user if cred else None


def _digits(value):
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def _conflict(model, field, value, exclude_user_id=None, label=""):
    """Return a human message if ``value`` is already used by another record."""
    if not value:
        return None
    qs = model.objects.filter(**{field: value})
    if exclude_user_id:
        qs = qs.exclude(user_id=exclude_user_id)
    obj = qs.first()
    if obj is None:
        return None
    who = ""
    for attr in ("bss_worker_id", "contractor_code", "employer_code"):
        who = getattr(obj, attr, "") or ""
        if who:
            break
    if not who:
        who = getattr(obj, "name", "") or getattr(obj, "contractor_name", "") or getattr(obj, "employer_name", "") or ""
    return f"{label} already registered{(' to ' + str(who)) if who else ''}."


def _template_response(entity):
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    title, columns = ENTITIES[entity]
    wb = Workbook()
    ws = wb.active
    ws.title = title
    fill = PatternFill("solid", fgColor="0A192F")
    for col, (label, _key) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        ws.column_dimensions[cell.column_letter].width = max(14, len(label) + 4)
    info = wb.create_sheet("Instructions")
    for i, line in enumerate([
        f"BSS Admin — Bulk register {title}", "",
        "Fill one row per record and save as .xlsx, then upload with the password.",
        "Mobile is the login id (10 digits). Existing mobiles are updated, not duplicated.",
        "Relations sheet links Employer Mobile ↔ Contractor Mobile ↔ Worker Mobile + Site.",
    ], start=1):
        info.cell(row=i, column=1, value=line)
    info.column_dimensions["A"].width = 110
    out = io.BytesIO()
    wb.save(out)
    resp = HttpResponse(out.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="bss-bulk-{entity}.xlsx"'
    return resp


# ── per-entity processors ────────────────────────────────────────────────────
def _process_employers(rows):
    from bss_core.models import User, UserCredential, EmployerProfile
    from api import bss_ids
    import re
    results = []
    seen_mobile = set()
    for r in rows:
        mobile = _mobile(r.get("mobile"))
        name = str(r.get("name") or "").strip()
        cin = str(r.get("cin") or "").strip().upper()
        gstin = str(r.get("gstin") or "").strip().upper()
        pan = str(r.get("pan") or "").strip().upper()
        if not mobile:
            results.append({"row": name or "(no name)", "status": "failed", "detail": "Mobile is required."}); continue
        if mobile in seen_mobile:
            results.append({"row": name or mobile, "status": "failed", "detail": "Duplicate mobile within the file."}); continue
        if gstin and not re.match(r"^\d{2}[A-Z]{5}\d{4}[A-Z]\d[Z][A-Z\d]$", gstin):
            results.append({"row": name or mobile, "status": "failed", "detail": "GSTIN format is invalid."}); continue
        if pan and not re.match(r"^[A-Z]{5}\d{4}[A-Z]$", pan):
            results.append({"row": name or mobile, "status": "failed", "detail": "PAN format is invalid (AAAAA9999A)."}); continue
        user = _find_user_by_mobile(mobile)
        if user is not None and user.role != "employer":
            results.append({"row": mobile, "status": "failed", "detail": f"Mobile belongs to a {user.role}."}); continue
        exclude = user.id if user else None
        conflict = (_conflict(EmployerProfile, "cin", cin, exclude, "CIN")
                    or _conflict(EmployerProfile, "gstin", gstin, exclude, "GSTIN")
                    or _conflict(EmployerProfile, "pan_no", pan, exclude, "PAN"))
        if conflict:
            results.append({"row": name or mobile, "status": "failed", "detail": conflict}); continue
        seen_mobile.add(mobile)
        if user is None:
            user = User.objects.create_user(username=f"employer_{mobile}", role="employer",
                                            account_active=True, is_verified=True, is_active=True)
            UserCredential.objects.create(user=user, identity_type="mobile", identity_value=mobile,
                                          is_primary=True, verified=True)
        profile, _ = EmployerProfile.objects.get_or_create(user=user)
        for field, key in (("employer_name", "name"), ("email", "email"), ("entity_class", "entity_type"),
                           ("cin", "cin"), ("gstin", "gstin"), ("pan_no", "pan"),
                           ("address", "address"), ("district", "district"), ("city", "city"),
                           ("state", "state"), ("pincode", "pincode"), ("contact_person", "contact_person")):
            if hasattr(profile, field):
                val = str(r.get(key) or "").strip()
                if val:
                    setattr(profile, field, val[:255])
        if name:
            profile.employer_name = name[:200]
        if not profile.employer_code:
            profile.employer_code = bss_ids.new_bss_employer_code()
        profile.save()
        results.append({"row": name or mobile, "status": "ok", "detail": f"employer_{mobile}"})
    return results


def _process_contractors(rows):
    from bss_core.models import User, UserCredential, ContractorProfile
    from api import bss_ids
    import re
    results = []
    seen_mobile = set()
    for r in rows:
        mobile = _mobile(r.get("mobile"))
        name = str(r.get("name") or "").strip()
        cin = str(r.get("cin") or "").strip().upper()
        gstin = str(r.get("gstin") or "").strip().upper()
        pan = str(r.get("pan") or "").strip().upper()
        clra = str(r.get("clra_licence_no") or "").strip().upper()
        if not mobile:
            results.append({"row": name or "(no name)", "status": "failed", "detail": "Mobile is required."}); continue
        if mobile in seen_mobile:
            results.append({"row": name or mobile, "status": "failed", "detail": "Duplicate mobile within the file."}); continue
        if gstin and not re.match(r"^\d{2}[A-Z]{5}\d{4}[A-Z]\d[Z][A-Z\d]$", gstin):
            results.append({"row": name or mobile, "status": "failed", "detail": "GSTIN format is invalid."}); continue
        if pan and not re.match(r"^[A-Z]{5}\d{4}[A-Z]$", pan):
            results.append({"row": name or mobile, "status": "failed", "detail": "PAN format is invalid (AAAAA9999A)."}); continue
        user = _find_user_by_mobile(mobile)
        if user is not None and user.role != "contractor":
            results.append({"row": mobile, "status": "failed", "detail": f"Mobile belongs to a {user.role}."}); continue
        exclude = user.id if user else None
        conflict = (_conflict(ContractorProfile, "cin", cin, exclude, "CIN")
                    or _conflict(ContractorProfile, "gstin", gstin, exclude, "GSTIN")
                    or _conflict(ContractorProfile, "pan_no", pan, exclude, "PAN")
                    or _conflict(ContractorProfile, "clra_licence_no", clra, exclude, "CLRA licence"))
        if conflict:
            results.append({"row": name or mobile, "status": "failed", "detail": conflict}); continue
        seen_mobile.add(mobile)
        if user is None:
            user = User.objects.create_user(username=f"contractor_{mobile}", role="contractor",
                                            account_active=True, is_verified=True, is_active=True)
            UserCredential.objects.create(user=user, identity_type="mobile", identity_value=mobile,
                                          is_primary=True, verified=True)
        profile, _ = ContractorProfile.objects.get_or_create(user=user)
        # Creating the profile stamps registered_at and consent_date together.
        _record_contractor_membership(
            profile, user,
            source="Admin Bulk Register", registered_by="Admin")
        for field, key in (("contractor_name", "name"), ("trade_name", "trade_name"),
                           ("constitution", "constitution"), ("cin", "cin"), ("gstin", "gstin"),
                           ("pan_no", "pan"), ("clra_licence_no", "clra_licence_no"),
                           ("operating_states", "operating_states"), ("address", "address"),
                           ("district", "district"), ("city", "city"), ("state", "state")):
            if hasattr(profile, field):
                val = str(r.get(key) or "").strip()
                if val:
                    setattr(profile, field, val[:255])
        lvt = _as_date(r.get("licence_valid_to"))
        if lvt and hasattr(profile, "licence_valid_to"):
            profile.licence_valid_to = lvt
        if name:
            profile.contractor_name = name[:200]
        if not profile.contractor_code:
            profile.contractor_code = bss_ids.new_bss_contractor_code()
        profile.save()
        results.append({"row": name or mobile, "status": "ok", "detail": f"contractor_{mobile}"})
    return results


def _process_workers(rows):
    from bss_core.models import (User, UserCredential, WorkerIdentity, WorkerSocialSecurity, WorkerOther)
    from api import bss_ids
    import re
    import uuid
    results = []
    seen = {"mobile": set(), "aadhaar": set(), "pan": set(), "uan": set(), "esic": set()}
    for r in rows:
        mobile = _mobile(r.get("mobile"))
        name = str(r.get("name") or "").strip()
        aadhaar = _digits(r.get("aadhaar"))
        pan = str(r.get("pan") or "").strip().upper()
        uan = _digits(r.get("uan"))
        esic = str(r.get("esic") or "").strip().upper()
        label = name or mobile or "(no name)"
        if not mobile:
            results.append({"row": label, "status": "failed", "detail": "Mobile is required."}); continue
        if aadhaar and len(aadhaar) != 12:
            results.append({"row": label, "status": "failed", "detail": "Aadhaar must be 12 digits."}); continue
        if aadhaar and aadhaar in seen["aadhaar"]:
            results.append({"row": label, "status": "failed", "detail": "Duplicate Aadhaar within the file."}); continue
        if pan and not re.match(r"^[A-Z]{5}\d{4}[A-Z]$", pan):
            results.append({"row": label, "status": "failed", "detail": "PAN format is invalid (AAAAA9999A)."}); continue
        if pan and pan in seen["pan"]:
            results.append({"row": label, "status": "failed", "detail": "Duplicate PAN within the file."}); continue
        if uan and uan in seen["uan"]:
            results.append({"row": label, "status": "failed", "detail": "Duplicate UAN within the file."}); continue
        if esic and esic in seen["esic"]:
            results.append({"row": label, "status": "failed", "detail": "Duplicate ESIC IP within the file."}); continue
        if mobile in seen["mobile"]:
            results.append({"row": label, "status": "failed", "detail": "Duplicate mobile within the file."}); continue

        user = _find_user_by_mobile(mobile)
        if user is not None and user.role != "worker":
            results.append({"row": mobile, "status": "failed", "detail": f"Mobile belongs to a {user.role}."}); continue
        exclude = user.id if user else None
        conflict = (_conflict(WorkerIdentity, "aadhaar", aadhaar, exclude, "Aadhaar")
                    or _conflict(WorkerIdentity, "pan", pan, exclude, "PAN")
                    or _conflict(WorkerSocialSecurity, "epfo_uan", uan, exclude, "UAN")
                    or _conflict(WorkerSocialSecurity, "esic_ip_number", esic, exclude, "ESIC IP"))
        if conflict:
            results.append({"row": label, "status": "failed", "detail": conflict}); continue

        seen["mobile"].add(mobile)
        if aadhaar:
            seen["aadhaar"].add(aadhaar)
        if pan:
            seen["pan"].add(pan)
        if uan:
            seen["uan"].add(uan)
        if esic:
            seen["esic"].add(esic)

        if user is None:
            user = User.objects.create_user(username=f"worker_{mobile}", role="worker",
                                            account_active=True, is_verified=True, is_active=True)
            UserCredential.objects.create(user=user, identity_type="mobile", identity_value=mobile,
                                          is_primary=True, verified=True)
        ident, _ = WorkerIdentity.objects.get_or_create(user=user)
        if not ident.bss_worker_id:
            ident.bss_worker_id = bss_ids.new_bss_worker_id()
        if not ident.qr_token:
            ident.qr_token = uuid.uuid4().hex + uuid.uuid4().hex[:8]
        if name:
            ident.name = name[:200]
        ident.mobile_number = mobile
        for field, key in (("father_name", "father_name"), ("aadhaar", "aadhaar"),
                           ("pan", "pan"), ("gender", "gender")):
            val = str(r.get(key) or "").strip()
            if val and hasattr(ident, field):
                setattr(ident, field, val[:120])
        dob = _as_date(r.get("dob"))
        if dob:
            ident.date_of_birth = dob
        ident.save()
        ss, _ = WorkerSocialSecurity.objects.get_or_create(user=user)
        if uan:
            ss.epfo_uan = uan[:60]
        if esic:
            ss.esic_ip_number = esic[:60]
        ss.save()
        WorkerOther.objects.get_or_create(user=user)
        results.append({"row": label, "status": "ok", "detail": ident.bss_worker_id})
    return results


def _process_relations(rows):
    from datetime import datetime
    from django.utils import timezone
    from bss_core.models import (Site, ContractorEmployerBinding, ContractorSiteAllotment,
                                 WorkerContractorBinding)
    from api import bss_ids
    results = []
    for r in rows:
        emp_m = _mobile(r.get("employer_mobile"))
        con_m = _mobile(r.get("contractor_mobile"))
        wrk_m = _mobile(r.get("worker_mobile"))
        site_name = str(r.get("site_name") or "").strip()
        label = f"{emp_m or '?'} / {con_m or '?'} / {wrk_m or site_name or '?'}"
        employer = _find_user_by_mobile(emp_m)
        contractor = _find_user_by_mobile(con_m)
        if employer is None or employer.role != "employer":
            results.append({"row": label, "status": "failed", "detail": "Employer mobile not found."}); continue
        if contractor is None or contractor.role != "contractor":
            results.append({"row": label, "status": "failed", "detail": "Contractor mobile not found."}); continue
        # Site (optional).
        site = None
        if site_name:
            site = (Site.objects.filter(site_name__iexact=site_name, employer=employer).first()
                    or Site.objects.filter(site_code__iexact=str(r.get("site_code") or "").strip(),
                                           employer=employer).first())
            if site is None:
                site = Site(employer=employer, site_name=site_name[:200], state="")
                code = str(r.get("site_code") or "").strip()
                site.site_code = code or bss_ids.new_bss_site_code()
                site.save()
        # Employer ↔ Contractor binding.
        binding, _ = ContractorEmployerBinding.objects.get_or_create(
            employer=employer, contractor=contractor,
            defaults={"binding_status": "ACCEPTED", "consent_contractor": True,
                      "consent_employer": True, "accepted_at": timezone.now(), "source": "ADMIN_BULK"})
        if binding.binding_status != "ACCEPTED":
            binding.binding_status = "ACCEPTED"; binding.accepted_at = timezone.now(); binding.save()
        if site is not None:
            ContractorSiteAllotment.objects.get_or_create(
                binding=binding, site=site,
                defaults={"status": "APPROVED", "source": "ADMIN_BULK", "approved_at": timezone.now()})
        # Worker ↔ Contractor + Site binding.
        if wrk_m:
            worker = _find_user_by_mobile(wrk_m)
            if worker is None or worker.role != "worker":
                results.append({"row": label, "status": "failed", "detail": "Worker mobile not found."}); continue
            wb, _ = WorkerContractorBinding.objects.get_or_create(
                worker=worker, contractor=contractor, site=site,
                defaults={"binding_status": "ACCEPTED", "initiated_by": "ADMIN_BULK",
                          "consent_worker": True, "consent_contractor": True,
                          "designation": str(r.get("designation") or "")[:120],
                          "accepted_at": timezone.now()})
            if wb.binding_status != "ACCEPTED":
                wb.binding_status = "ACCEPTED"; wb.accepted_at = timezone.now(); wb.save()
        results.append({"row": label, "status": "ok", "detail": "linked"})
    return results


_PROCESSORS = {"employers": _process_employers, "contractors": _process_contractors,
               "workers": _process_workers, "relations": _process_relations}


@api_view(["GET", "POST"])
@authentication_classes([BssSessionAuthentication])
@permission_classes([IsAuthenticated])
def admin_bulk_register(request, entity):
    """Password-protected bulk registration. GET = XLSX template; POST = upload."""
    entity = str(entity or "").strip().lower()
    if entity not in ENTITIES:
        return Response({"error": "Unknown entity."}, status=status.HTTP_404_NOT_FOUND)
    denied = _authorize(request)
    if denied is not None:
        return denied
    if request.method == "GET":
        return _template_response(entity)
    upload = request.FILES.get("file") or request.FILES.get("master")
    if upload is None:
        return Response({"error": "Upload the filled .xlsx file."},
                        status=status.HTTP_400_BAD_REQUEST)
    try:
        rows = _rows(upload, entity)
    except Exception as exc:  # noqa: BLE001
        return Response({"error": f"Could not read the file: {exc}"},
                        status=status.HTTP_400_BAD_REQUEST)
    if not rows:
        return Response({"error": "No data rows found in the file."},
                        status=status.HTTP_400_BAD_REQUEST)
    try:
        results = _PROCESSORS[entity](rows)
    except Exception as exc:  # noqa: BLE001
        return Response({"error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
    summary = {"total": len(results), "ok": 0, "failed": 0}
    for r in results:
        summary[r["status"]] = summary.get(r["status"], 0) + 1
    return Response({"data": {"results": results, "summary": summary}}, status=status.HTTP_200_OK)


__all__ = ["admin_bulk_register"]
