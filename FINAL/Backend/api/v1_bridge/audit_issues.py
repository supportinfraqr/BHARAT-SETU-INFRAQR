"""Build an Excel (.xlsx) listing the **pointwise AI-audit problems** for a
document / register / submission, so it can be downloaded.

`build_issues_xlsx(title, meta, problems)` where ``problems`` is a list of
``{"point": str, "detail": str}``. `problems_from_audit(audit)` pulls the
pointwise problems out of a stored Level-1 audit result.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEAD_FILL = PatternFill("solid", fgColor="073B67")


def problems_from_audit(audit):
    """Pointwise problems from a stored audit dict (problems + failed checks)."""
    audit = audit or {}
    out, seen = [], set()

    def _add(point, detail=""):
        point = str(point or "").strip()
        if not point:
            return
        key = point.lower()
        if key in seen:
            return
        seen.add(key)
        out.append({"point": point[:900], "detail": str(detail or "")[:900]})

    for p in (audit.get("problems") or []):
        _add(p)
    for c in (audit.get("checkpointResults") or []):
        if isinstance(c, dict) and c.get("passed") is False:
            _add(f"Condition not met: {c.get('label')}", c.get("detail") or "")
    for x in (audit.get("issues") or []):
        _add(x)
    for x in (audit.get("missing") or []):
        _add(f"Missing: {x}")
    for x in (audit.get("findings") or []):
        _add(x)
    return out


def problems_from_evidence(ev):
    """Pointwise problems from a stored SubmissionEvidence row."""
    out, seen = [], set()

    def _add(point, detail=""):
        point = str(point or "").strip()
        if not point:
            return
        if point.lower() in seen:
            return
        seen.add(point.lower())
        out.append({"point": point[:900], "detail": str(detail or "")[:900]})

    for c in (ev.checkpoint_results or []):
        if isinstance(c, dict) and c.get("passed") is False:
            _add(f"Condition not met: {c.get('label')}", c.get("detail") or "")
    for x in (ev.validation_findings or []):
        _add(x)
    return out


def build_issues_xlsx(title, meta, problems):
    """Bytes of an .xlsx with a meta header and the pointwise problem table."""
    wb = Workbook()
    ws = wb.active
    ws.title = "AI Audit Issues"

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="073B67")
    r = 3
    for key, value in meta:
        ws.cell(row=r, column=1, value=str(key)).font = Font(bold=True)
        ws.cell(row=r, column=2, value=str(value or ""))
        r += 1
    r += 1

    for col, header in enumerate(["#", "Problem / Failed condition", "Detail"], start=1):
        cell = ws.cell(row=r, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    r += 1

    if not problems:
        ws.cell(row=r, column=2, value="No issues found.").font = Font(italic=True, color="15803D")
    else:
        for i, p in enumerate(problems, start=1):
            ws.cell(row=r, column=1, value=i)
            ws.cell(row=r, column=2, value=p.get("point") or "")
            ws.cell(row=r, column=3, value=p.get("detail") or "")
            r += 1

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 52
    for row in ws.iter_rows(min_row=1, max_row=r, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
