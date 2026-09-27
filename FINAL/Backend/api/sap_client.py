"""SAP → ContractorProfile integration.

Two ways to pull a contractor's master data from SAP:

* **Live** — SAP OData (S/4HANA / SAP Gateway), e.g. the Business Partner API
  ``API_BUSINESS_PARTNER``. Configured via env::

      BSS_SAP_BASE_URL=https://sap.example.com/sap/opu/odata/sap/API_BUSINESS_PARTNER
      BSS_SAP_ENTITY=A_BusinessPartner
      BSS_SAP_ID_FIELD=BusinessPartner
      BSS_SAP_USERNAME=...        # Basic auth (or)
      BSS_SAP_TOKEN=...           # Bearer token
      BSS_SAP_VERIFY_TLS=1

  ``fetch_business_partner(identifier)`` returns the first matching record as a
  flat ``{field: value}`` dict.

* **File** — an SAP export (CSV/XLSX) is parsed and the row matching the
  contractor code is used (``parse_sap_export``). This works with any SAP
  (IDoc/CREM/table export) with no connectivity.

``map_sap_record`` maps the SAP record onto :class:`ContractorProfile` fields
using ``SAP_FIELD_MAP`` (candidate SAP field aliases → profile field).
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import urllib.parse
import urllib.request
from datetime import date, datetime

# Profile field → human label, and the SAP field aliases that feed it.
SAP_FIELD_MAP: list[tuple[str, str, tuple[str, ...]]] = [
    ("contractor_code", "Contractor code", ("businesspartner", "businesspartnerid", "lifnr", "vendor", "vendorno", "contractorcode")),
    ("contractor_name", "Contractor name", ("organizationbpname1", "name1", "name", "contractorname", "businesspartnerfullname", "displayname", "vendorname")),
    ("trade_name", "Trade / brand name", ("searchterm1", "sortl", "tradename", "brandname")),
    ("constitution", "Constitution", ("legalform", "legalformname", "constitution")),
    ("entity_type", "Entity type", ("businesspartnercategory", "entitytype")),
    ("cin", "CIN", ("bpidentificationnumber", "cin", "companyregistrationnumber")),
    ("roc", "ROC", ("rocno", "roc", "registrarofcompanies")),
    ("registration_number", "Registration number", ("registrationnumber", "regnof", "regno", "companyregno")),
    ("nic_code", "NIC code", ("niccode", "nic", "industrynic")),
    ("members_count", "Members count", ("memberscount", "members", "numberofmembers")),
    ("authorized_capital", "Authorised capital", ("authorizedcapital", "authorisedcapital")),
    ("paid_up_capital", "Paid-up capital", ("paidupcapital", "paidupcapital")),
    ("incorporation_date", "Incorporation date", ("incorporationdate", "dateofincorporation")),
    ("agm_date", "AGM date", ("agmdate", "lastagmdate")),
    ("balance_sheet_date", "Balance-sheet date", ("balancesheetdate", "balancesheetasofdate")),
    ("gstin", "GSTIN", ("taxnumber1", "gstin", "gstno", "gstnumber")),
    ("pan_no", "PAN", ("taxnumber2", "panno", "pannumber", "bpidentificationnumber2")),
    ("udyam_no", "Udyam / MSME", ("udyamno", "udyam", "msmeno", "msmenumber")),
    ("labour_dept_reg_no", "Labour dept. reg no.", ("labourdeptregno", "labourdepotregno")),
    ("clra_licence_no", "CLRA licence no.", ("clralicenceno", "clralicence", "licenceno", "licensenumber")),
    ("bocw_reg_no", "BOCW reg no.", ("bocwregno", "bocwreg")),
    ("licence_valid_from", "Licence valid from", ("licencevalidfrom", "validfrom", "licencefromdate")),
    ("licence_valid_to", "Licence valid to", ("licencevalidto", "validto", "licencevalidtodate", "licenceexpiry")),
    ("operating_states", "Operating states", ("operatingstates", "statesofoperation")),
    ("contact_person", "Contact person", ("contactperson", "contactpersonname", "contactpersonfullname")),
    ("mobile", "Mobile", ("telephone1", "telf1", "mobileno", "mobile", "contactnumber", "phone")),
    ("email", "Email", ("emailaddress", "smtpaddr", "email", "emailid")),
    ("address", "Address", ("streetname", "stras", "address", "street")),
    ("district", "District / city", ("city", "ort01", "cityname", "district")),
    ("state", "State", ("region", "regio", "state", "statename")),
]

DATE_FIELDS = {"incorporation_date", "agm_date", "balance_sheet_date",
               "licence_valid_from", "licence_valid_to"}
DECIMAL_FIELDS = {"authorized_capital", "paid_up_capital"}
INT_FIELDS = {"members_count"}

# Fields SAP must never write. consent_date is stamped the moment the BSS
# profile is created and kept equal to registered_at, so an external system
# must not be able to set or move it. Excluded here so it cannot arrive even
# through a per-employer fieldMapping override.
_NON_SAP_FIELDS = {"consent_date"}


def _norm(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _parse_date(value):
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _coerce(field: str, value):
    if value is None or str(value).strip() == "":
        return None
    if field in DATE_FIELDS:
        return _parse_date(value)
    if field in DECIMAL_FIELDS:
        try:
            return float(str(value).replace(",", "").strip())
        except ValueError:
            return None
    if field in INT_FIELDS:
        try:
            return int(float(str(value).strip()))
        except ValueError:
            return None
    return str(value).strip()[:255]


# Extra aliases for SAP Business One (Service Layer) BusinessPartner fields.
B1_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "contractor_code": ("cardcode",),
    "contractor_name": ("cardname",),
    "trade_name": ("foreignname", "cardforeignname"),
    "category_name": ("groupname",),
    "gstin": ("lictradnum", "gstin", "gstno", "gstnumber"),
    "pan_no": ("federaltaxid", "panno", "additionalid", "pannumber"),
    "contact_person": ("contactperson", "contactpersonname"),
    "mobile": ("phone1", "cellular", "phone2", "mobilephone"),
    "email": ("emailaddress", "e_mail", "email"),
    "address": ("street", "mailaddress", "addressname"),
    "district": ("city",),
    "state": ("state",),
    "clra_licence_no": ("clralicence", "clralicenceno"),
}


def map_sap_record(record: dict, overrides: dict | None = None) -> dict:
    """Map a raw SAP record (any header casing) to ContractorProfile fields.

    ``overrides`` is a per-employer map ``{bss_field: [sap field, …]}`` applied
    ahead of the built-in aliases (so an employer can pin e.g. ``address`` ←
    ``VendorAddress``)."""
    flat = {_norm(k): v for k, v in (record or {}).items() if k}
    extra: dict[str, tuple[str, ...]] = {}
    for bss_field, sap_fields in (overrides or {}).items():
        if bss_field in _NON_SAP_FIELDS:
            continue
        names = [sap_fields] if isinstance(sap_fields, str) else list(sap_fields or [])
        extra[bss_field] = tuple(_norm(n) for n in names)
    out: dict = {}
    for field, _label, aliases in SAP_FIELD_MAP:
        if field in _NON_SAP_FIELDS:
            continue
        for alias in extra.get(field, ()) + aliases + B1_FIELD_ALIASES.get(field, ()):
            if alias in flat and str(flat[alias]).strip() != "":
                coerced = _coerce(field, flat[alias])
                if coerced is not None:
                    out[field] = coerced
                break
    # Overrides for fields outside the built-in map.
    for field, aliases in extra.items():
        if field in out:
            continue
        for alias in aliases:
            if alias in flat and str(flat[alias]).strip() != "":
                coerced = _coerce(field, flat[alias])
                if coerced is not None:
                    out[field] = coerced
                break
    return out


def field_map_payload() -> list[dict]:
    """The mapping table for the UI (profile field · label · SAP aliases)."""
    return [{"field": field, "label": label, "sapFields": list(aliases)}
            for field, label, aliases in SAP_FIELD_MAP]


# ── file (SAP export) ────────────────────────────────────────────────────────
def _rows_from_file(file_obj) -> list[dict]:
    raw = file_obj.read()
    name = (getattr(file_obj, "name", "") or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        rows = [tuple(r) for r in csv.reader(io.StringIO(text))]
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for values in rows[1:]:
        row = {header[i]: values[i] for i in range(min(len(header), len(values))) if header[i]}
        if any(str(v or "").strip() for v in row.values()):
            out.append(row)
    return out


def parse_sap_export(file_obj, identifier: str = "") -> dict:
    """Pick the record from an SAP export matching ``identifier`` (code/CIN/PAN),
    else the first row."""
    rows = _rows_from_file(file_obj)
    if not rows:
        return {}
    ident = _norm(identifier)
    if ident:
        for row in rows:
            if ident in {_norm(v) for v in row.values()}:
                return row
    return rows[0]


# ── live OData ───────────────────────────────────────────────────────────────
def sap_config() -> dict:
    return {
        # "b1" = SAP Business One Service Layer, "odata" = generic SAP Gateway OData.
        "mode": (os.environ.get("BSS_SAP_MODE") or "odata").strip().lower(),
        "baseUrl": (os.environ.get("BSS_SAP_BASE_URL") or "").rstrip("/"),
        "companyDb": os.environ.get("BSS_SAP_COMPANYDB") or "",
        "entity": os.environ.get("BSS_SAP_ENTITY") or "A_BusinessPartner",
        "idField": os.environ.get("BSS_SAP_ID_FIELD") or "BusinessPartner",
        "username": os.environ.get("BSS_SAP_USERNAME") or "",
        "password": os.environ.get("BSS_SAP_PASSWORD") or "",
        "token": os.environ.get("BSS_SAP_TOKEN") or "",
        "verifyTls": (os.environ.get("BSS_SAP_VERIFY_TLS", "1") not in ("0", "false", "False")),
    }


def is_configured() -> bool:
    cfg = sap_config()
    return bool(cfg["baseUrl"])


def resolve_config(employer=None) -> dict:
    """Merge the employer's stored SAP config over the global env defaults."""
    base = sap_config()
    if employer is None:
        return base
    from bss_core.models import EmployerSapConfig
    from api import crypto
    row = EmployerSapConfig.objects.filter(employer=employer, is_active=True).first()
    if row is None:
        return base
    return {
        "mode": row.mode or base["mode"],
        "baseUrl": (row.base_url or base["baseUrl"]).rstrip("/"),
        "companyDb": row.company_db or base["companyDb"],
        "entity": row.entity or base["entity"],
        "idField": row.id_field or base["idField"],
        "username": row.username or base["username"],
        "password": crypto.decrypt(row.password_enc) or base.get("password", ""),
        "token": crypto.decrypt(row.token_enc) or base.get("token", ""),
        "authType": row.auth_type or "session",
        "extraLogin": row.extra_login or {},
        "verifyTls": bool(row.verify_tls),
        "bpFilter": row.bp_filter or "",
        "fieldMapping": row.field_mapping or {},
        "allowlistIps": row.allowlist_ips or "",
        "configured": bool(row.base_url or base["baseUrl"]),
    }


def fetch_business_partner(identifier: str, cfg: dict | None = None) -> dict:
    """Fetch one Business Partner from SAP OData by ``identifier``. Raises
    ``RuntimeError`` when SAP isn't configured or the request fails."""
    cfg = cfg or sap_config()
    if not cfg["baseUrl"]:
        raise RuntimeError("SAP is not configured (set BSS_SAP_BASE_URL).")
    query = urllib.parse.urlencode({"$filter": f"{cfg['idField']} eq '{identifier}'", "$top": 1, "$format": "json"})
    url = f"{cfg['baseUrl']}/{cfg['entity']}?{query}"
    request = urllib.request.Request(url)
    request.add_header("Accept", "application/json")
    if cfg.get("token"):
        request.add_header("Authorization", f"Bearer {cfg['token']}")
    elif cfg.get("username"):
        import base64
        password = cfg.get("password") or ""
        token = base64.b64encode(f"{cfg['username']}:{password}".encode()).decode()
        request.add_header("Authorization", f"Basic {token}")
    import ssl
    ctx = None if cfg["verifyTls"] else ssl._create_unverified_context()
    try:
        with urllib.request.urlopen(request, timeout=30, context=ctx) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"SAP request failed: {exc}") from exc
    results = (payload.get("d", {}) or {}).get("results") or payload.get("value") or []
    if not results:
        raise RuntimeError(f"No SAP record found for '{identifier}'.")
    return results[0]


# ── SAP Business One Service Layer ───────────────────────────────────────────
def _ssl_context(verify: bool):
    import ssl
    return None if verify else ssl._create_unverified_context()


def b1_login(cfg: dict | None = None) -> dict:
    """Log in to the B1 Service Layer and return the session cookies
    (``B1SESSION``, ``RouteID``). With ``authType=bearer`` + a token, returns a
    sentinel ``{"__bearer__": token}`` instead of logging in."""
    cfg = cfg or sap_config()
    if not cfg["baseUrl"]:
        raise RuntimeError("SAP is not configured (set BSS_SAP_BASE_URL).")
    if (cfg.get("authType") or "session") == "bearer" and cfg.get("token"):
        return {"__bearer__": cfg["token"]}
    login_body = {
        "CompanyDB": cfg["companyDb"],
        "UserName": cfg["username"],
        "Password": cfg.get("password") or "",
    }
    login_body.update(cfg.get("extraLogin") or {})
    body = json.dumps(login_body).encode("utf-8")
    request = urllib.request.Request(f"{cfg['baseUrl']}/Login", data=body, method="POST")
    request.add_header("Content-Type", "application/json")
    request.add_header("Accept", "application/json")
    try:
        with urllib.request.urlopen(request, timeout=30,
                                    context=_ssl_context(cfg["verifyTls"])) as resp:
            cookies = resp.headers.get_all("Set-Cookie") or []
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"SAP B1 login failed: {exc}") from exc
    jar: dict = {}
    for raw in cookies:
        head = raw.split(";", 1)[0]
        if "=" in head:
            name, _, value = head.partition("=")
            jar[name.strip()] = value.strip()
    if "B1SESSION" not in jar:
        raise RuntimeError("SAP B1 login did not return a B1SESSION.")
    return jar


def _flatten_b1_bp(bp: dict) -> dict:
    """Flatten a B1 BusinessPartner (nested addresses/contacts) to a flat dict."""
    flat = {k: v for k, v in (bp or {}).items() if not isinstance(v, (list, dict))}
    addresses = bp.get("BPAddresses") or []
    if addresses:
        chosen = next((a for a in addresses
                       if str(a.get("AddressType", "")).lower().endswith("billto")), addresses[0])
        for key in ("Street", "City", "ZipCode", "State", "Country", "AddressName"):
            if chosen.get(key) and not flat.get(key):
                flat[key] = chosen[key]
    contacts = bp.get("ContactEmployees") or []
    if contacts:
        contact = contacts[0]
        if contact.get("Name"):
            flat.setdefault("ContactPersonName", contact["Name"])
        for src, dst in (("Phone1", "Phone1"), ("MobilePhone", "Cellular"), ("E_Mail", "E_Mail")):
            if contact.get(src) and not flat.get(dst):
                flat[dst] = contact[src]
    return flat


def b1_get_business_partner(card_code: str, cfg: dict | None = None) -> dict:
    """Fetch one Business Partner from the B1 Service Layer by ``CardCode``."""
    cfg = cfg or sap_config()
    if not cfg["baseUrl"]:
        raise RuntimeError("SAP is not configured (set BSS_SAP_BASE_URL).")
    jar = b1_login(cfg)
    cookie = ""
    bearer = jar.get("__bearer__")
    url = f"{cfg['baseUrl']}/BusinessPartners('{urllib.parse.quote(card_code)}')"
    request = urllib.request.Request(url)
    request.add_header("Accept", "application/json")
    if bearer:
        request.add_header("Authorization", f"Bearer {bearer}")
    else:
        cookie = "; ".join(f"{k}={v}" for k, v in jar.items())
        request.add_header("Cookie", cookie)
    try:
        with urllib.request.urlopen(request, timeout=30,
                                    context=_ssl_context(cfg["verifyTls"])) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
        # Best-effort logout (session mode only).
        if cookie:
            try:
                logout = urllib.request.Request(f"{cfg['baseUrl']}/Logout", data=b"{}", method="POST")
                logout.add_header("Cookie", cookie)
                urllib.request.urlopen(logout, timeout=10, context=_ssl_context(cfg["verifyTls"])).close()
            except Exception:  # noqa: BLE001
                pass
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"SAP B1 request failed: {exc}") from exc
    if not payload:
        raise RuntimeError(f"No SAP B1 Business Partner found for '{card_code}'.")
    return _flatten_b1_bp(payload)


def fetch_record(identifier: str, cfg: dict | None = None) -> dict:
    """Fetch a SAP record by identifier using the configured mode (b1|odata)."""
    cfg = cfg or sap_config()
    if cfg.get("mode") == "b1":
        return b1_get_business_partner(identifier, cfg)
    return fetch_business_partner(identifier, cfg)


def allowlist_ips(cfg: dict | None = None) -> list[str]:
    """Outbound IP(s) the employer must allow-list in/at their SAP so the BSS
    server can reach the SAP API. Prefers the employer's stored
    ``allowlistIps``, else ``BSS_SAP_ALLOWLIST_IPS``, else detected."""
    if cfg:
        stored = cfg.get("allowlistIps")
        if isinstance(stored, str) and stored.strip():
            return [x.strip() for x in re.split(r"[,\s]+", stored) if x.strip()]
    raw = (os.environ.get("BSS_SAP_ALLOWLIST_IPS") or "").strip()
    ips = [x.strip() for x in re.split(r"[,\s]+", raw) if x.strip()]
    if ips:
        return ips
    try:
        with urllib.request.urlopen("https://api.ipify.org", timeout=6) as resp:
            public_ip = resp.read().decode("utf-8").strip()
        if public_ip:
            return [public_ip]
    except Exception:  # noqa: BLE001
        pass
    import socket
    try:
        return sorted({socket.gethostbyname(socket.gethostname())})
    except Exception:  # noqa: BLE001
        return []
