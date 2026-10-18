"""Reads Agent1_AuditableFields_v3.xlsx — the audit catalogue the agent must cover, row by row.

Sheet "Possible"     : S.No | Criteria | Check Area | Item Compared | Rulebook Ref | Remarks | Act
Sheet "Not Possible" : S.No | Criteria | Check Area | Item | Rulebook Ref | Reason Not Possible | Act

Every row becomes a CheckSpec, and every bullet / line of a Possible row's Remarks becomes a
sub-check that must receive its own verdict. The engine guarantees one result per row and per
sub-check, so editing the workbook changes the audit scope without code edits.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class CheckSpec:
    sheet: str              # "Possible" | "Not Possible"
    sno: int
    criteria: str           # audit area (report section), e.g. "Wage Register"
    area: str               # check area (Possible: category such as "Wages & statutory calculation")
    item: str
    ref: str
    remarks: str = ""       # Possible: client remarks; Not Possible: reason
    act: str = ""
    subchecks: list[str] = field(default_factory=list)
    ref_codes: list[str] = field(default_factory=list)

    @property
    def cid(self) -> str:
        return f"{'P' if self.sheet == 'Possible' else 'NP'}-{self.sno:03d}"

    @property
    def area_code(self) -> str:
        m = re.match(r"(A\d{2})", self.area) or re.match(r"(A\d{2})", self.ref or "")
        return m.group(1) if m else ""


def _clean(v) -> str:
    return "" if v is None else re.sub(r"[ \t]+", " ", str(v)).strip()


def split_remarks(text: str) -> list[str]:
    """Remarks -> independent sub-checks. Bullets ('•') or line breaks separate them; a single
    sentence stays one sub-check. Very wide spaces inside a cell are line breaks typed as spaces."""
    t = re.sub(r" {6,}", "\n", str(text or ""))
    parts = re.split(r"\n|•", t)
    out = [re.sub(r"\s+", " ", p).strip(" -–;") for p in parts]
    return [p for p in out if len(p) > 3]


def load_catalogue(path: Path) -> list[CheckSpec]:
    wb = load_workbook(path, read_only=False, data_only=True)
    specs: list[CheckSpec] = []
    for sheet in ("Possible", "Not Possible"):
        if sheet not in wb.sheetnames:
            raise ValueError(f"{path.name}: sheet '{sheet}' is missing")
        ws = wb[sheet]
        # merged cells (e.g. one Remarks spanning two rows) apply to every row they cover
        merged = {}
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged[(r, c)] = v
        rows = []
        for r in range(1, ws.max_row + 1):
            rows.append([merged.get((r, c), ws.cell(r, c).value) for c in range(1, ws.max_column + 1)])
        header = [_clean(h).lower() for h in rows[0]]
        col = {name: i for i, name in enumerate(header)}

        def get(r, *names, raw=False):
            for n in names:
                if n in col and col[n] < len(r):
                    v = r[col[n]]
                    return ("" if v is None else str(v)) if raw else _clean(v)
            return ""

        for r in rows[1:]:
            sno = r[col["s.no"]] if r else None
            if sno in (None, ""):
                continue
            remarks = get(r, "remarks", "reason not possible")
            ref = get(r, "rulebook ref")
            specs.append(CheckSpec(
                sheet=sheet, sno=int(sno), criteria=get(r, "criteria") or get(r, "check area"),
                area=get(r, "check area"), item=get(r, "item compared", "item"), ref=ref, remarks=remarks,
                act=get(r, "act"), subchecks=split_remarks(get(r, "remarks", raw=True)) if sheet == "Possible" else [],
                ref_codes=re.findall(r"A\d{2}\.\d+|§\d{2}", ref)))
    ids = [s.cid for s in specs]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        raise ValueError(f"{path.name}: duplicate S.No {sorted(dup)}")
    return specs
