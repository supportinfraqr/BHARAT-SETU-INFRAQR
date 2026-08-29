#!/usr/bin/env python3
"""
Generate seed_rbac.sql + seed_qr.sql + seed_lookup.sql from the source workbook.

seed_rbac.sql  source: "Worker form & Access" sheet  (field-level permission matrix)
seed_qr.sql    source: "QR Role Summary" sheet       (per-role QR access profile)
seed_lookup.sql source: "Dropdowns" sheet            (enumerable pick-list values)
seed_worker_category.sql source: "Worker Categories" sheet (W01-W20 coded master taxonomy)
seed_skill_master.sql    source: "Skill Master" sheet      (S01-S45 coded skill taxonomy)
seed_document_master.sql source: "Documents" sheet         (14 worker document types)

RBAC matrix mapping:
Maps the 10 matrix role columns onto the 9 auth roles. The "Contractor Supervisor"
and "Contractor Bound Worker" columns both map to the single `contractor` role,
differentiated by the `binding_scope` column (unbound / bound). All other roles map
1:1.

Access token normalization:
    "Yes" / "yes"        -> view  (public-safe display)
    "Verified name only" -> view  (reveal_type = verified_only)
    "View"               -> view
    "Masked"             -> masked
    "View/Masked"        -> view_masked
    "Restricted"         -> restricted
    "No" / "" / "-"      -> no    (deny by default)
    "Full"               -> full
    "Optional"           -> optional (public display only when worker opted in)

Each (role x resource x field) combination emits one INSERT row. For `contractor`
two rows are emitted (unbound + bound). The `resource` is derived from the matrix
`Section` value (snake_cased) so it lines up with the data tables.

SPECIAL_RESOURCE_FIELDS: fields that live in dedicated role/site tables rather than
`worker_employment`, and are therefore re-scoped to their owning resource so the RBAC
rows still govern them even though they are no longer stored inline on the worker row.

LOOKUP parsing (Dropdowns sheet): the sheet lists pick-list values as vertical columns.
The first 25 columns (0-24) are the primary dropdowns; columns 25-46 duplicate most of
them plus add Employment Relationship & Current Employment Status; a final stacked
group (column 47) holds sub-group dropdowns (site induction, gate passes, gig,
self-employment, season). Values are de-duplicated and one row is emitted per
(category, value) into lookup_value.

Run:
    python3 generate_seed_rbac.py
Writes: seed_rbac.sql, seed_qr.sql, seed_lookup.sql, seed_worker_category.sql, seed_skill_master.sql, seed_document_master.sql
"""
import re
import openpyxl

SRC = "/home/yasser/BSS_Project/FINAL/Requirements/BSS_Final_worker_form_and_permission.xlsx"
SHEET = "Worker form & Access"
QR_SHEET = "QR Role Summary"
LOOKUP_SHEET = "Dropdowns"
WC_SHEET = "Worker Categories"
SM_SHEET = "Skill Master"
DOC_SHEET = "Documents"
OUT = "seed_rbac.sql"
OUT_QR = "seed_qr.sql"
OUT_LOOKUP = "seed_lookup.sql"
OUT_WC = "seed_worker_category.sql"
OUT_SM = "seed_skill_master.sql"
OUT_DOC = "seed_document_master.sql"

# Dropdowns sheet -> lookup_value category mapping (column index -> category).
# Block 1 (cols 0-24): the primary form dropdowns.
LOOKUP_COLUMNS = {
    0:  "worker_category",
    1:  "sector",
    2:  "industry",
    3:  "occupation",
    4:  "skill",
    5:  "gender",
    6:  "marital_status",
    7:  "disability_verification_status",
    8:  "migration_type",
    9:  "shift",
    10: "platform_type",
    11: "vehicle_type",
    12: "vehicle_ownership",
    13: "qualification",
    14: "skill_level",
    15: "digital_literacy",
    16: "identity_verification_status",
    17: "bank_verification_status",
    18: "nominee_relationship",
    19: "medical_examination_status",
    20: "medical_fitness_status",
    21: "insurance_status",
    22: "language",
    23: "communication_channel",
    24: "registration_source",
    # Block 2 unique additions (cols 28, 29) not already covered above.
    28: "employment_relationship",
    29: "current_employment_status",
}

# Column 47 (stacked sub-group dropdowns): each group begins with a header row of the
# form "N. <Category Label>" followed by its values until the next header. Map the
# header label's canonical category to a snake_case lookup category.
LOOKUP_47_CATEGORIES = {
    "migrant_worker": "migrant_worker",
    "site_induction_status": "site_induction_status",
    "temporary_gate_pass_status": "gate_pass_status",
    "permanent_gate_pass_status": "gate_pass_status",
    "gig_active_platform_type": "platform_type",
    "gig_active_vehicle_ownership": "gig_vehicle_ownership",
    "self_employment_business_type": "se_business_type",
    "self_employment_business_registration_status": "se_registration_status",
    "season_job_work_type": "season_job_type",
}

# index of the 10 matrix columns after [Sl.No, Section, Field, Dropdown, Type, Owner, Mandatory]
COLS = [
    ("public",              None),                    # Role-Public
    ("contractor",          "unbound"),               # Role-Contractor Supervisor
    ("contractor",          "bound"),                 # Role-Contractor Bound Worker
    ("employer",            None),                    # Role-Employer Bound Worker
    ("government_inspector",None),                    # Role-Government Authorised Inspector
    ("government_admin",    None),                    # Role-Government Admin
    ("worker",              None),                    # Role-Worker Own Profile
    ("it_task_force",       None),                    # Role-BSS-Tech Team
    ("bss_super_admin",     None),                    # Role-BSS super Admin
    ("bss_staff",           None),                    # Role-BSS-Staff
]

# "QR Role Summary" role label -> auth role code.
# Only roles that scan a worker's QR in the field are modelled: the public/unknown
# scanner, the worker's own profile, the contractor (unbound supervisor + bound worker)
# and the employer (bound). Back-office roles (government, IT / BSS staff, super admin)
# are NOT QR-scanner roles and are intentionally excluded from qr_role_access.
QR_ROLE_MAP = {
    "Public / Unknown Scanner": "public",
    "Worker – Own Profile":     "worker",
    "Contractor / Supervisor – Before Binding": "contractor",  # -> unbound
    "Contractor – Bound Worker": "contractor",                  # -> bound
    "Employer – Bound Worker":  "employer",
}

# Which QR labels map to which binding_scope (only meaningful for contractor).
QR_BINDING_SCOPE = {
    "Public / Unknown Scanner": None,
    "Worker – Own Profile":     None,
    "Contractor / Supervisor – Before Binding": "unbound",
    "Contractor – Bound Worker": "bound",
    "Employer – Bound Worker":  None,
}


def normalize(token):
    """Return (access, reveal_type_or_None)."""
    low = (token or "").strip().lower()
    if low in ("yes", "yes "):
        return "view", None
    if low == "verified name only":
        return "view", "verified_only"
    if low == "view":
        return "view", None
    if low == "masked":
        return "masked", None
    if low in ("view/masked", "view / masked", "viewmasked"):
        return "view_masked", None
    if low == "restricted":
        return "restricted", None
    if low == "full":
        return "full", None
    if low == "optional":
        return "optional", None
    if low in ("", "-", "no"):
        return "no", None
    if "view" in low:
        return "view", None
    return "no", None


def resource_from_section(section):
    s = re.sub(r"[^A-Za-z0-9]+", "_", (section or "").strip().lower()).strip("_")
    return s or "worker_profile"


def field_key(field):
    f = re.sub(r"[^A-Za-z0-9]+", "_", (field or "").strip().lower()).strip("_")
    if not f:
        f = "field"
    return f


# Fields whose data lives in a dedicated role/site table (not worker_employment).
# Map field_key -> owning resource. Used to re-scope RBAC rows.
SPECIAL_RESOURCE_FIELDS = {
    # site (owns site_name / site_code / site_location / state fields)
    "site_name":    "site",
    "site_code":    "site",
    "site_location":"site",
    # state appears once (row 125 = Employment/Work State). It refers to the site's state.
    "state":        "site",
    # contractor profile
    "contractor_name":    "contractor_profiles",
    "contractor_code":    "contractor_profiles",
    "contractor_email":   "contractor_profiles",
    "contractor_state":   "contractor_profiles",
    "contractor_address": "contractor_profiles",
    "contractor_lin":     "contractor_profiles",
    # employer profile
    "employer_name":      "employer_profiles",
    "employer_code":      "employer_profiles",
    "employer_email":     "employer_profiles",
    "employer_email_id":  "employer_profiles",
    "employer_state":     "employer_profiles",
    "employer_address":   "employer_profiles",
    "employer_industry":  "employer_profiles",
    "employer_lin":       "employer_profiles",
}


def resource_for(section, field_key):
    """Return the RBAC resource for a field, honouring the special re-scoping."""
    if field_key in SPECIAL_RESOURCE_FIELDS:
        return SPECIAL_RESOURCE_FIELDS[field_key]
    return resource_from_section(section)


def boolish(token):
    """True for yes/true/1, False for no/false/0; None otherwise."""
    low = (token or "").strip().lower()
    if low in ("yes", "true", "1"):
        return True
    if low in ("no", "false", "0"):
        return False
    return None


def write_qr_seed(wb):
    """Generate seed_qr.sql from the 'QR Role Summary' sheet."""
    ws = wb[QR_SHEET]
    rows = []
    for r, values in enumerate(ws.iter_rows(values_only=True), 1):
        if r == 1:
            continue
        cells = ["" if c is None else str(c).strip() for c in values]
        if not any(cells):
            continue
        role_label = cells[0]
        role = QR_ROLE_MAP.get(role_label)
        if role is None:
            continue
        bscope = QR_BINDING_SCOPE.get(role_label)
        grant = cells[1]
        default_access = cells[2]
        # 'Can view sensitive identifiers?' column: "No" / "Yes, to own data" / "Masked..." / ...
        sensitive_raw = cells[3]
        # A clear "yes" means sensitive identifiers can be viewed in routine access.
        sensitive = boolish(sensitive_raw)
        if sensitive is None:
            sensitive = False
        optional_raw = cells[4]
        # 'Recommended scope' mentioning 'photo if consented' / 'opted in' implies opt-in gating.
        optional = False
        if optional_raw and re.search(r"consent|opt", optional_raw, re.I):
            optional = True
        rows.append((role, bscope, grant, default_access, sensitive, optional, optional_raw))

    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- Bharat Shram Setu (BSS) — QR Role Summary seed")
    lines.append(f"-- Generated from '{QR_SHEET}' sheet of the source workbook: {len(rows)} roles")
    lines.append("-- Only QR-scanner roles: public, worker, contractor (unbound+bound), employer.")
    lines.append("-- Regenerate with: python3 generate_seed_rbac.py")
    lines.append("-- =============================================================================")
    lines.append("")
    lines.append("BEGIN;")
    lines.append("")
    lines.append(
        "INSERT INTO qr_role_access (role, binding_scope, grant_mechanism, default_access, "
        "sensitive_allowed, optional_identifiers, recommended_scope) VALUES"
    )
    lines.append("")
    for i, (role, bscope, grant, da, sensitive, optional, scope) in enumerate(rows):
        comma = "," if i < len(rows) - 1 else ";"
        s = ("'" + scope.replace("'", "''") + "'") if scope else "NULL"
        bv = ("'" + bscope + "'") if bscope else "NULL"
        lines.append(
            f"    ('{role}', {bv}, '{grant.replace(chr(39), chr(39)+chr(39))}', "
            f"'{da.replace(chr(39), chr(39)+chr(39))}', "
            f"{'TRUE' if sensitive else 'FALSE'}, {'TRUE' if optional else 'FALSE'}, {s}){comma}"
        )
    lines.append("")
    lines.append("COMMIT;")

    with open(OUT_QR, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Wrote {OUT_QR}: {len(rows)} qr_role_access rows")


def snake(s):
    return re.sub(r"[^A-Za-z0-9]+", "_", (s or "").strip().lower()).strip("_")


def col_values(data, idx):
    """Extract non-empty column values preserving order and de-duplicating."""
    seen = []
    for r in data[1:]:
        v = r[idx] if idx < len(r) else ""
        v = v.strip()
        if v and v not in seen:
            seen.append(v)
    return seen


def read_lookup(data):
    """Return {category: [values,...]} parsed from the Dropdowns sheet data."""
    categories = {}

    def add(cat, vals):
        cat = snake(cat)
        if not cat:
            return
        cur = categories.setdefault(cat, [])
        for v in vals:
            if v not in cur:
                cur.append(v)

    # Block 1 + block 2 unique columns.
    for idx, cat in LOOKUP_COLUMNS.items():
        vals = col_values(data, idx)
        if vals:
            add(cat, vals)

    # Column 47: stacked sub-groups. Each group begins with a header of the form
    # "N. <Label>" followed by values until the next header. The first group header
    # ('1. Migrant Worker') lives in the header row itself, so seed `group` from it.

    def resolve_group_label(label):
        """Map a stripped label to a lookup category via the alias table."""
        ln = snake(label)
        for k, v in LOOKUP_47_CATEGORIES.items():
            if snake(k) == ln:
                return v
        return LOOKUP_47_CATEGORIES.get(ln, ln)

    first_hdr = (data[0][47] if len(data[0]) > 47 else "") or ""
    group = None
    if first_hdr:
        m0 = re.match(r"^\d+\.\s*(.+)$", first_hdr)
        if m0:
            group = resolve_group_label(m0.group(1))

    for r in data[1:]:
        cell = (r[47] if len(r) > 47 else "") or ""
        cell_s = snake(cell)
        # Match header patterns: "1. Migrant Worker", "2. Site Induction Status" ...
        m = re.match(r"^\d+\.\s*(.+)$", cell)
        if m:
            group = resolve_group_label(m.group(1))
            continue
        if group is None:
            continue
        if cell_s:
            add(group, [cell])

    return categories


def write_lookup_seed(wb):
    """Generate seed_lookup.sql from the 'Dropdowns' sheet."""
    ws = wb[LOOKUP_SHEET]
    data = []
    for values in ws.iter_rows(values_only=True):
        cells = [("" if c is None else str(c).strip()) for c in values]
        data.append(cells)

    categories = read_lookup(data)

    # Emit ordered rows grouped by category.
    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- Bharat Shram Setu (BSS) — Dropdown / reference-data seed")
    lines.append(f"-- Generated from '{LOOKUP_SHEET}' sheet of the source workbook")
    total = sum(len(v) for v in categories.values())
    lines.append(f"-- {len(categories)} categories, {total} lookup values")
    lines.append("-- Regenerate with: python3 generate_seed_rbac.py")
    lines.append("-- =============================================================================")
    lines.append("")
    lines.append("BEGIN;")
    lines.append("")
    lines.append("INSERT INTO lookup_value (category, value, sort_order) VALUES")
    rows = []
    for cat, vals in categories.items():
        for i, v in enumerate(vals):
            rows.append((cat, v, i))
    for i, (cat, v, order) in enumerate(rows):
        comma = "," if i < len(rows) - 1 else ";"
        vesc = v.replace("'", "''")
        lines.append(f"    ('{cat}', '{vesc}', {order}){comma}")
    lines.append("")
    lines.append("COMMIT;")

    with open(OUT_LOOKUP, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Wrote {OUT_LOOKUP}: {len(categories)} categories, {total} values")


def write_worker_category_seed(wb):
    """Generate seed_worker_category.sql from the 'Worker Categories' sheet."""
    ws = wb[WC_SHEET]
    rows = []
    for r, values in enumerate(ws.iter_rows(values_only=True), 1):
        if r == 1:
            continue  # header
        cells = ["" if c is None else str(c).strip() for c in values]
        if not any(cells):
            continue
        code = cells[0] if len(cells) > 0 else ""
        name = cells[1] if len(cells) > 1 else ""
        if not code or not name:
            continue
        definition = cells[2] if len(cells) > 2 else ""
        workplaces = cells[3] if len(cells) > 3 else ""
        rel = cells[4] if len(cells) > 4 else ""
        focus = cells[5] if len(cells) > 5 else ""
        rows.append((code, name, definition, workplaces, rel, focus, len(rows)))

    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- Bharat Shram Setu (BSS) — Worker Category Master seed")
    lines.append(f"-- Generated from '{WC_SHEET}' sheet of the source workbook: {len(rows)} categories")
    lines.append("-- Coded taxonomy (W01-W20) for worker_classification.worker_category.")
    lines.append("-- Regenerate with: python3 generate_seed_rbac.py")
    lines.append("-- =============================================================================")
    lines.append("")
    lines.append("BEGIN;")
    lines.append("")
    lines.append(
        "INSERT INTO worker_category_master (code, name, definition, typical_workplaces, "
        "primary_relationship, data_focus, sort_order) VALUES"
    )
    lines.append("")
    for i, (code, name, definition, workplaces, rel, focus, order) in enumerate(rows):
        comma = "," if i < len(rows) - 1 else ";"

        def esc(v):
            return "NULL" if not v else "'" + v.replace("'", "''") + "'"

        lines.append(
            f"    ('{code}', '{name.replace(chr(39), chr(39)+chr(39))}', "
            f"{esc(definition)}, {esc(workplaces)}, {esc(rel)}, {esc(focus)}, {order}){comma}"
        )
    lines.append("")
    lines.append("COMMIT;")

    with open(OUT_WC, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Wrote {OUT_WC}: {len(rows)} worker_category_master rows")


def write_skill_master_seed(wb):
    """Generate seed_skill_master.sql from the 'Skill Master' sheet."""
    ws = wb[SM_SHEET]
    rows = []
    for r, values in enumerate(ws.iter_rows(values_only=True), 1):
        if r == 1:
            continue  # header
        cells = ["" if c is None else str(c).strip() for c in values]
        if not any(cells):
            continue
        code = cells[0] if len(cells) > 0 else ""
        sector = cells[1] if len(cells) > 1 else ""
        if not code or not sector:
            continue
        family = cells[2] if len(cells) > 2 else ""
        occupation = cells[3] if len(cells) > 3 else ""
        level = cells[4] if len(cells) > 4 else ""
        industry = cells[5] if len(cells) > 5 else ""
        desc = cells[6] if len(cells) > 6 else ""
        rows.append((code, sector, family, occupation, level, industry, desc, len(rows)))

    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- Bharat Shram Setu (BSS) — Skill Master seed")
    lines.append(f"-- Generated from '{SM_SHEET}' sheet of the source workbook: {len(rows)} skills")
    lines.append("-- Coded taxonomy (S01-S45) for skill reference data.")
    lines.append("-- Regenerate with: python3 generate_seed_rbac.py")
    lines.append("-- =============================================================================")
    lines.append("")
    lines.append("BEGIN;")
    lines.append("")
    lines.append(
        "INSERT INTO skill_master (code, sector, family, occupation, typical_level, "
        "industry_use, description, sort_order) VALUES"
    )
    lines.append("")
    for i, (code, sector, family, occupation, level, industry, desc, order) in enumerate(rows):
        comma = "," if i < len(rows) - 1 else ";"

        def esc(v):
            return "NULL" if not v else "'" + v.replace("'", "''") + "'"

        lines.append(
            f"    ('{code}', {esc(sector)}, {esc(family)}, {esc(occupation)}, {esc(level)}, "
            f"{esc(industry)}, {esc(desc)}, {order}){comma}"
        )
    lines.append("")
    lines.append("COMMIT;")

    with open(OUT_SM, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Wrote {OUT_SM}: {len(rows)} skill_master rows")


def write_document_master_seed(wb):
    """Generate seed_document_master.sql from the 'Documents' sheet."""
    ws = wb[DOC_SHEET]
    rows = []
    for r, values in enumerate(ws.iter_rows(values_only=True), 1):
        if r == 1:
            continue  # header
        cells = ["" if c is None else str(c).strip() for c in values]
        if not any(cells):
            continue
        name = cells[2] if len(cells) > 2 else ""
        if not name:
            continue
        mandatory = (cells[5].lower() == "yes") if len(cells) > 5 else False
        owner = cells[4] if len(cells) > 4 else "Worker"
        uploaded_by = "system" if "system" in owner.lower() else "worker"
        rows.append((name, mandatory, uploaded_by, len(rows)))

    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- Bharat Shram Setu (BSS) — Worker Document Master seed")
    lines.append(f"-- Generated from '{DOC_SHEET}' sheet of the source workbook: {len(rows)} document types")
    lines.append("-- Master list of worker document uploads (DOC-01 ...).")
    lines.append("-- Regenerate with: python3 generate_seed_rbac.py")
    lines.append("-- =============================================================================")
    lines.append("")
    lines.append("BEGIN;")
    lines.append("")
    lines.append(
        "INSERT INTO worker_document_master (code, document_name, is_mandatory, uploaded_by, sort_order) VALUES"
    )
    lines.append("")
    for i, (name, mandatory, uploaded_by, order) in enumerate(rows):
        comma = "," if i < len(rows) - 1 else ";"
        code = f"DOC-{i+1:02d}"
        mand = "TRUE" if mandatory else "FALSE"
        lines.append(
            f"    ('{code}', '{name.replace(chr(39), chr(39)+chr(39))}', {mand}, "
            f"'{uploaded_by}', {order}){comma}"
        )
    lines.append("")
    lines.append("COMMIT;")

    with open(OUT_DOC, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Wrote {OUT_DOC}: {len(rows)} worker_document_master rows")


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    write_qr_seed(wb)
    write_lookup_seed(wb)
    write_worker_category_seed(wb)
    write_skill_master_seed(wb)
    write_document_master_seed(wb)
    ws = wb[SHEET]

    rows = []
    for r, values in enumerate(ws.iter_rows(values_only=True), 1):
        if r == 1:
            continue
        cells = ["" if c is None else str(c).strip() for c in values]
        if not any(cells):
            continue
        sl, section, field = cells[0], cells[1], cells[2]
        perms = cells[7:17]  # the 10 permission columns (indices 7..16)
        if len(perms) < 10:
            continue
        rows.append((sl, section, field, perms))

    insert_rows = []
    # (access, reveal) per generated (role, binding_scope)
    for sl, section, field, perms in rows:
        fkey = field_key(field)
        resource = resource_for(section, fkey)
        for idx, (role, bscope) in enumerate(COLS):
            access, reveal = normalize(perms[idx])
            # For non-contractor roles binding_scope stays NULL; for contractor set it.
            bval = bscope if role == "contractor" else None
            insert_rows.append((role, resource, fkey, access, reveal, bval))

    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- Bharat Shram Setu (BSS) — FINAL RBAC permission seed")
    lines.append(f"-- Generated from '{SHEET}' sheet of the source workbook: {len(rows)} fields")
    lines.append(f"-- {len(insert_rows)} rbac_permissions rows (contractor emits bound + unbound rows)")
    lines.append("-- Regenerate with: python3 generate_seed_rbac.py")
    lines.append("-- =============================================================================")
    lines.append("")
    lines.append("BEGIN;")
    lines.append("")
    lines.append("INSERT INTO rbac_permissions (role, resource, field, access, reveal_type, binding_scope) VALUES")
    lines.append("")
    for i, (role, resource, fkey, access, reveal, bval) in enumerate(insert_rows):
        rv = ("'" + reveal + "'") if reveal else "NULL"
        bv = ("'" + bval + "'") if bval else "NULL"
        comma = "," if i < len(insert_rows) - 1 else ";"
        lines.append(
            f"    ('{role}', '{resource}', '{fkey}', '{access}', {rv}, {bv}){comma}"
        )
    lines.append("")
    lines.append("-- Note: this seed intentionally sets every role x resource x field from the matrix.")
    lines.append("-- 'no' covers blank / '-' cells (deny-by-default).")
    lines.append("")
    lines.append("COMMIT;")

    with open(OUT, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Wrote {OUT}: {len(insert_rows)} permission rows across {len(rows)} fields")


if __name__ == "__main__":
    main()
